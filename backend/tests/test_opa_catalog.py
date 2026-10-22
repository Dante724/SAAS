"""Team A (Orderry parity) — catalog and public page.

A1  catalog PUTs write only the fields that were sent (no silent resets).
A2  Service.bookable drives the public appointment page.
A3  Product MPN: CRUD, uniqueness, search, scan, posting import, Excel, v1 API, {Product MPN}.
A4  Label templates choose the barcode source field and a symbology per field (EAN-13/UPC-A
    validated, Code128/QR kept).
Each new or changed endpoint has a permission-denied and a cross-tenant check.
"""
import io
import uuid

import pytest
from sqlalchemy import select

from backend.services.codes import (
    ean13_modules,
    gtin_check_digit,
    label_barcode,
    normalize_gtin,
)


async def _register(client, name="Co"):
    email = f"opa-{uuid.uuid4().hex[:8]}@t.dev"
    r = await client.post("/api/auth/register", json={
        "org_name": name, "name": "Owner", "email": email, "password": "pw-test-12345"})
    assert r.status_code == 201, r.text
    body = r.json()
    H = {"Authorization": f"Bearer {body['access_token']}"}
    me = (await client.get("/api/auth/me", headers=H)).json()
    return H, body["org_id"], me["id"]


async def _teammate(client, H, *deny):
    """A teammate whose role has every permission except `deny`."""
    r = await client.post("/api/settings/rbac/roles", json={
        "name": f"No {deny[0]} {uuid.uuid4().hex[:4]}", "permissions": {c: False for c in deny}}, headers=H)
    assert r.status_code in (200, 201), r.text
    email = f"tm-{uuid.uuid4().hex[:8]}@t.dev"
    r = await client.post("/api/staff/employees", json={
        "name": "Tech", "email": email, "password": "pw-test-12345", "role_id": r.json()["id"]}, headers=H)
    assert r.status_code == 201, r.text
    tok = await client.post("/api/auth/token", data={"username": email, "password": "pw-test-12345"})
    return r.json()["id"], {"Authorization": f"Bearer {tok.json()['access_token']}"}


async def _product(client, H, **kw):
    body = {"name": f"Part {uuid.uuid4().hex[:6]}", "sell_price": 10, **kw}
    r = await client.post("/api/catalog/products", json=body, headers=H)
    assert r.status_code == 201, r.text
    return r.json()


# ── A1: service save keeps every field it wasn't sent ─────────────────────────
FULL_SERVICE = {
    "name": "Screen labour", "description": "Teardown", "is_labor": True, "unit": "h",
    "price": 120, "cost_price": 40, "duration_min": 90, "flat_rate": True,
    "warranty_days": 90, "barcode": "SVC-1", "bookable": False,
    "required_resource_types": ["bay", "Lift"],
}


async def test_service_partial_put_keeps_unsent_fields(client):
    H, _, me = await _register(client)
    s = (await client.post("/api/catalog/services", json={**FULL_SERVICE, "allowed_employee_ids": [me]},
                           headers=H)).json()
    # the drawer's old body: none of warranty/barcode/bookable/booking rules
    r = await client.put(f"/api/catalog/services/{s['id']}", json={
        "name": "Screen labour v2", "price": 130}, headers=H)
    assert r.status_code == 200, r.text
    u = r.json()
    assert u["name"] == "Screen labour v2" and u["price"] == 130
    assert u["warranty_days"] == 90 and u["barcode"] == "SVC-1" and u["bookable"] is False
    assert u["allowed_employee_ids"] == [me]
    assert u["required_resource_types"] == ["bay", "Lift"]
    assert u["description"] == "Teardown" and u["unit"] == "h" and u["duration_min"] == 90
    assert u["is_labor"] is True and u["flat_rate"] is True and u["cost_price"] == 40


async def test_service_full_put_round_trips_and_clears(client):
    H, _, me = await _register(client)
    s = (await client.post("/api/catalog/services", json=FULL_SERVICE, headers=H)).json()
    r = await client.put(f"/api/catalog/services/{s['id']}", json={
        **FULL_SERVICE, "warranty_days": 0, "barcode": "", "bookable": True,
        "allowed_employee_ids": [me, me], "required_resource_types": [" room ", "room"]}, headers=H)
    u = r.json()
    assert u["warranty_days"] == 0 and u["barcode"] is None and u["bookable"] is True
    assert u["allowed_employee_ids"] == [me]               # de-duplicated
    assert u["required_resource_types"] == ["room"]         # trimmed + de-duplicated


async def test_service_rejects_another_orgs_employee(client):
    H, _, _ = await _register(client)
    _, _, other_emp = await _register(client, "Other")
    r = await client.post("/api/catalog/services", json={"name": "X", "allowed_employee_ids": [other_emp]},
                          headers=H)
    assert r.status_code == 400
    s = (await client.post("/api/catalog/services", json={"name": "X"}, headers=H)).json()
    r = await client.put(f"/api/catalog/services/{s['id']}", json={
        "name": "X", "allowed_employee_ids": [other_emp]}, headers=H)
    assert r.status_code == 400


async def test_service_put_needs_permission_and_own_org(client):
    H, _, _ = await _register(client)
    s = (await client.post("/api/catalog/services", json={"name": "Diag"}, headers=H)).json()
    _, T = await _teammate(client, H, "can_edit_labors")
    assert (await client.put(f"/api/catalog/services/{s['id']}", json={"name": "Diag"},
                             headers=T)).status_code == 403
    H2, _, _ = await _register(client, "Other")
    assert (await client.put(f"/api/catalog/services/{s['id']}", json={"name": "Diag"},
                             headers=H2)).status_code == 404


async def test_product_partial_put_keeps_unsent_fields(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, barcode="X1", barcode_type="qr", weight=1.5, dimensions="1x2x3",
                       image_data="data:image/png;base64,AA==", supplier_warranty_days=30,
                       mpn="MPN-1", cost_price=4, serialized=True)
    r = await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "sell_price": 12},
                         headers=H)
    assert r.status_code == 200, r.text     # a serialized product saves without resending `serialized`
    u = r.json()
    assert u["sell_price"] == 12
    assert u["barcode_type"] == "qr" and u["weight"] == 1.5 and u["dimensions"] == "1x2x3"
    assert u["image_data"] == "data:image/png;base64,AA==" and u["supplier_warranty_days"] == 30
    assert u["mpn"] == "MPN-1" and u["cost_price"] == 4 and u["serialized"] is True
    # explicitly turning serialization off is still refused
    r = await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "serialized": False},
                         headers=H)
    assert r.status_code == 409


async def test_product_put_keeps_legacy_category_name(client, db):
    """A product with only a category *name* (no managed category) keeps it on a partial save."""
    from backend.models import Part
    H, _, _ = await _register(client)
    p = await _product(client, H)
    row = (await db.execute(select(Part).where(Part.id == p["id"]))).scalar_one()
    row.category = "Screens"
    await db.commit()
    await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "sell_price": 9}, headers=H)
    await db.refresh(row)
    assert row.category == "Screens"


async def test_bundle_partial_put_keeps_components_and_category(client):
    H, _, _ = await _register(client)
    svc = (await client.post("/api/catalog/services", json={"name": "Labour", "price": 50}, headers=H)).json()
    cat = (await client.post("/api/catalog/categories", json={"entity_type": "bundle", "name": "Kits"},
                             headers=H)).json()
    b = (await client.post("/api/bundles", json={
        "name": "Kit", "price_mode": "fixed", "price": 80, "category_id": cat["id"],
        "components": [{"ref_type": "service", "ref_id": svc["id"], "qty": 2}]}, headers=H)).json()
    r = await client.put(f"/api/bundles/{b['id']}", json={"name": "Kit 2"}, headers=H)
    assert r.status_code == 200, r.text
    u = r.json()
    assert u["name"] == "Kit 2" and u["price_mode"] == "fixed" and u["price"] == 80
    assert u["category_id"] == cat["id"] and len(u["components"]) == 1 and u["components"][0]["qty"] == 2
    r = await client.put(f"/api/bundles/{b['id']}", json={"name": "Kit 2", "components": []}, headers=H)
    assert r.status_code == 400


async def test_tax_partial_put_keeps_unsent_fields(client):
    H, _, _ = await _register(client)
    t = (await client.post("/api/catalog/taxes", json={
        "name": "VAT", "rate": 20, "code": "S", "name_for_document": "VAT 20%"}, headers=H)).json()
    u = (await client.put(f"/api/catalog/taxes/{t['id']}", json={"name": "VAT std"}, headers=H)).json()
    assert u["name"] == "VAT std" and u["rate"] == 20 and u["code"] == "S"
    assert u["name_for_document"] == "VAT 20%"


# ── A2: bookable drives the public appointment page ───────────────────────────
async def test_unbookable_service_is_hidden_from_public_page(client, db):
    from backend.models import Organisation
    H, org_id, _ = await _register(client)
    on = (await client.post("/api/catalog/services", json={"name": "Visible svc"}, headers=H)).json()
    off = (await client.post("/api/catalog/services", json={"name": "Hidden svc"}, headers=H)).json()
    await client.put(f"/api/catalog/services/{off['id']}", json={"name": "Hidden svc", "bookable": False},
                     headers=H)
    org = (await db.execute(select(Organisation).where(Organisation.id == org_id))).scalar_one()
    org.appointment_enabled = True
    await db.commit()
    r = await client.get(f"/api/public/appointment/{org_id}")
    assert r.status_code == 200, r.text
    ids = {s["id"] for s in r.json()["services"]}
    assert on["id"] in ids and off["id"] not in ids


# ── A3: MPN ───────────────────────────────────────────────────────────────────
async def test_mpn_round_trip_and_uniqueness(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, mpn="  661-12345  ")
    assert p["mpn"] == "661-12345"
    got = (await client.get(f"/api/catalog/products/{p['id']}", headers=H)).json()["product"]
    assert got["mpn"] == "661-12345"
    # another product with the same MPN (any case) is refused, on create and on edit
    r = await client.post("/api/catalog/products", json={"name": "Dup", "mpn": "661-12345".lower()}, headers=H)
    assert r.status_code == 409 and p["name"] in r.json()["detail"]
    q = await _product(client, H)
    r = await client.put(f"/api/catalog/products/{q['id']}", json={"name": q["name"], "mpn": "661-12345"},
                         headers=H)
    assert r.status_code == 409
    # re-saving the same product with its own MPN is fine; blank clears it
    assert (await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "mpn": "661-12345"},
                             headers=H)).status_code == 200
    u = (await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "mpn": ""},
                          headers=H)).json()
    assert u["mpn"] is None
    # another org may use the same MPN
    H2, _, _ = await _register(client, "Other")
    await _product(client, H2, mpn="661-12345")


async def test_mpn_too_long_is_rejected(client):
    H, _, _ = await _register(client)
    r = await client.post("/api/catalog/products", json={"name": "Long", "mpn": "x" * 65}, headers=H)
    assert r.status_code == 422


async def test_mpn_search_everywhere_and_tenant_scoped(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, name="Battery pack", mpn="BAT-77Q")
    H2, _, _ = await _register(client, "Other")
    await _product(client, H2, name="Other battery", mpn="BAT-77Q-X")

    items = (await client.get("/api/catalog/products?search=77q", headers=H)).json()["items"]
    assert [i["id"] for i in items] == [p["id"]] and items[0]["mpn"] == "BAT-77Q"
    found = (await client.get("/api/catalog/search?q=BAT-77&type=product", headers=H)).json()["items"]
    assert [i["ref_id"] for i in found] == [p["id"]] and found[0]["mpn"] == "BAT-77Q"
    parts = (await client.get("/api/inventory/parts?search=bat-77", headers=H)).json()["items"]
    assert [i["id"] for i in parts] == [p["id"]] and parts[0]["mpn"] == "BAT-77Q"
    scan = (await client.get("/api/inventory/scan?code=BAT-77Q", headers=H)).json()
    assert scan["product"]["id"] == p["id"] and scan["matched_on"] == "mpn"
    # the other org's exact MPN doesn't resolve here
    assert (await client.get("/api/inventory/scan?code=BAT-77Q-X", headers=H)).status_code == 404


async def test_posting_preview_matches_and_commit_creates_with_mpn(client):
    H, _, _ = await _register(client)
    loc = (await client.get("/api/staff/locations", headers=H)).json()["items"][0]["id"]
    p = await _product(client, H, name="Screen 13", mpn="SCR-13-OEM")
    prev = (await client.post("/api/inventory/postings/preview", json={"rows": [
        {"mpn": "scr-13-oem", "qty": 2},          # MPN column
        {"code": "SCR-13-OEM", "qty": 1},         # supplier "code" that is really the MPN
        {"name": "Brand new part", "mpn": "NEW-1", "qty": 3},
    ]}, headers=H)).json()["rows"]
    assert prev[0]["match_status"] == "matched" and prev[0]["part_id"] == p["id"]
    assert prev[1]["part_id"] == p["id"]
    assert prev[2]["match_status"] == "new" and prev[2]["mpn"] == "NEW-1"
    r = await client.post("/api/inventory/postings/commit", json={"location_id": loc, "rows": [
        {"new_name": "Brand new part", "new_mpn": "NEW-1", "qty": 3, "purchase_price": 5}]}, headers=H)
    assert r.status_code == 201, r.text
    items = (await client.get("/api/catalog/products?search=NEW-1", headers=H)).json()["items"]
    assert len(items) == 1 and items[0]["mpn"] == "NEW-1"


async def test_mpn_excel_export_and_import(client):
    import openpyxl
    H, _, _ = await _register(client)
    p = await _product(client, H, name="Cam module", mpn="CAM-9")
    r = await client.get("/api/catalog/products/export", headers=H)
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert "MPN" in header
    row = [c.value for c in ws[2]]
    assert row[header.index("MPN")] == "CAM-9"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "MPN"])
    ws.append(["Cam module", "CAM-10"])
    ws.append(["Fresh part", "FR-1"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await client.post("/api/catalog/products/import", headers=H,
                          files={"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert (await client.get(f"/api/catalog/products/{p['id']}", headers=H)).json()["product"]["mpn"] == "CAM-10"
    assert (await client.get("/api/catalog/products?search=FR-1", headers=H)).json()["items"][0]["mpn"] == "FR-1"

    # an older sheet without the MPN column leaves MPNs alone
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Sell price"])
    ws.append(["Cam module", 99])
    buf = io.BytesIO()
    wb.save(buf)
    await client.post("/api/catalog/products/import", headers=H,
                      files={"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")})
    assert (await client.get(f"/api/catalog/products/{p['id']}", headers=H)).json()["product"]["mpn"] == "CAM-10"


async def test_mpn_needs_manage_inventory(client):
    H, _, _ = await _register(client)
    p = await _product(client, H)
    _, T = await _teammate(client, H, "can_manage_inventory")
    assert (await client.post("/api/catalog/products", json={"name": "N", "mpn": "M"}, headers=T)).status_code == 403
    assert (await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "mpn": "M"},
                             headers=T)).status_code == 403
    H2, _, _ = await _register(client, "Other")
    assert (await client.put(f"/api/catalog/products/{p['id']}", json={"name": p["name"], "mpn": "M"},
                             headers=H2)).status_code == 404


async def test_mpn_in_public_api(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, mpn="API-MPN")
    key = (await client.post("/api/api-keys", json={"name": "zap", "scopes": "*"}, headers=H)).json()
    key = key.get("key") or key.get("token") or key.get("plaintext")
    auth = {"Authorization": f"Bearer {key}"}
    one = (await client.get(f"/api/v1/products/{p['id']}", headers=auth)).json()["data"]
    assert one["mpn"] == "API-MPN"
    listed = (await client.get("/api/v1/products", headers=auth)).json()["data"]
    assert any(x["mpn"] == "API-MPN" for x in listed)


async def test_product_mpn_print_variable(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, mpn="PRN-MPN")
    d = (await client.get("/api/documents/templates", headers=H)).json()
    assert "{Product MPN}" in d["variables"]
    tpl = (await client.post("/api/documents/templates", json={
        "name": "Tag", "doc_type": "label", "page_size": "label", "page_w_mm": 50, "page_h_mm": 30,
        "body_html": "<p>M={Product MPN}</p>"}, headers=H)).json()
    r = await client.post("/api/documents/render-labels", json={"template_id": tpl["id"], "part_ids": [p["id"]]},
                          headers=H)
    assert r.status_code == 200, r.text
    assert "M=PRN-MPN" in r.json()["html"]


# ── A4: barcode source + symbology ────────────────────────────────────────────
def test_gtin_validation():
    assert gtin_check_digit("400638133393") == "1"
    assert normalize_gtin("400638133393", "ean13") == "4006381333931"   # check digit appended
    assert normalize_gtin("4006381333931", "ean13") == "4006381333931"
    assert normalize_gtin("03600029145", "upca") == "036000291452"
    for bad, kind in (("4006381333932", "ean13"), ("12345", "ean13"), ("ABC123456789", "ean13"),
                      ("036000291453", "upca")):
        with pytest.raises(ValueError):
            normalize_gtin(bad, kind)


def test_ean13_modules_match_reportlab():
    """Our SVG bars must be the same pattern reportlab prints on the PDF labels."""
    from reportlab.graphics.barcode import eanbc
    for code in ("4006381333931", "5901234123457", "0036000291452"):
        w = eanbc.Ean13BarcodeWidget(code[:12])
        b = [w._tail]
        w._encode_left(code, b.append)
        b.append(w._sep)
        b += [w._right[int(c)] for c in code[7:]]
        b.append(w._tail)
        assert "".join(b) == ean13_modules(code)
        assert len(ean13_modules(code)) == 95


def test_label_barcode_picks_field_and_type():
    class P:
        barcode, sku, code, mpn = None, "SKU-1", "C-1", "590123412345"
    assert label_barcode(P, "mpn", {"mpn": "ean13"}) == ("590123412345", "ean13")
    assert label_barcode(P, "barcode", {}) == (None, "code128")          # chosen field empty
    assert label_barcode(P, None, {"sku": "qr"}) == ("SKU-1", "qr")      # legacy: first filled


async def test_code_svg_renders_ean_and_rejects_bad(client):
    H, _, _ = await _register(client)
    r = await client.get("/api/catalog/code.svg?type=ean13&data=400638133393", headers=H)
    assert r.status_code == 200 and "4006381333931" in r.text
    r = await client.get("/api/catalog/code.svg?type=upca&data=036000291452", headers=H)
    assert r.status_code == 200
    assert (await client.get("/api/catalog/code.svg?type=ean13&data=123", headers=H)).status_code == 400
    assert (await client.get("/api/catalog/code.svg?type=pdf417&data=1", headers=H)).status_code == 400


async def test_document_label_template_encodes_chosen_field(client):
    H, _, _ = await _register(client)
    good = await _product(client, H, name="Good", mpn="590123412345", sku="SKU-G")
    tpl = (await client.post("/api/documents/templates", json={
        "name": "MPN tag", "doc_type": "label", "page_size": "label", "page_w_mm": 50, "page_h_mm": 30,
        "body_html": "<div>{Barcode}</div>", "barcode_field": "mpn",
        "barcode_types": {"mpn": "ean13", "sku": "code128"}}, headers=H)).json()
    listed = [t for t in (await client.get("/api/documents/templates", headers=H)).json()["items"]
              if t["id"] == tpl["id"]][0]
    assert listed["barcode_field"] == "mpn" and listed["barcode_types"] == {"mpn": "ean13", "sku": "code128"}
    r = await client.post("/api/documents/render-labels", json={"template_id": tpl["id"], "part_ids": [good["id"]]},
                          headers=H)
    assert r.status_code == 200, r.text
    assert "5901234123457" in r.json()["html"]          # EAN-13 caption with its check digit

    bad = await _product(client, H, name="Bad", mpn="NOT-A-GTIN")
    r = await client.post("/api/documents/render-labels", json={"template_id": tpl["id"], "part_ids": [bad["id"]]},
                          headers=H)
    assert r.status_code == 422 and "Bad" in r.json()["detail"]

    # an editor save that doesn't send the barcode settings keeps them
    await client.put(f"/api/documents/templates/{tpl['id']}", json={
        "name": "MPN tag 2", "doc_type": "label", "body_html": "<div>{Barcode}</div>"}, headers=H)
    again = [t for t in (await client.get("/api/documents/templates", headers=H)).json()["items"]
             if t["id"] == tpl["id"]][0]
    assert again["barcode_field"] == "mpn" and again["barcode_types"]["mpn"] == "ean13"


async def test_document_template_barcode_options_validated_and_gated(client):
    H, _, _ = await _register(client)
    r = await client.post("/api/documents/templates", json={
        "name": "T", "doc_type": "label", "barcode_field": "colour"}, headers=H)
    assert r.status_code == 422
    r = await client.post("/api/documents/templates", json={
        "name": "T", "doc_type": "label", "barcode_types": {"mpn": "pdf417"}}, headers=H)
    assert r.status_code == 422
    _, T = await _teammate(client, H, "can_manage_settings")
    r = await client.post("/api/documents/templates", json={
        "name": "T", "doc_type": "label", "barcode_field": "mpn"}, headers=T)
    assert r.status_code == 403
    tpl = (await client.post("/api/documents/templates", json={"name": "T", "doc_type": "label"},
                             headers=H)).json()
    H2, _, _ = await _register(client, "Other")
    r = await client.put(f"/api/documents/templates/{tpl['id']}", json={
        "name": "T", "doc_type": "label", "barcode_field": "mpn"}, headers=H2)
    assert r.status_code == 404


async def test_print_label_template_barcode_field(client):
    H, _, _ = await _register(client)
    good = await _product(client, H, name="Good", mpn="400638133393", sku="S1")
    bad = await _product(client, H, name="Bad", mpn="xx", sku="S2")
    r = await client.post("/api/print/templates", json={
        "name": "MPN", "barcode_field": "mpn", "barcode_types": {"mpn": "ean13"},
        "fields": ["name", "mpn", "barcode"]}, headers=H)
    assert r.status_code == 201, r.text
    tpl = r.json()
    assert tpl["barcode_field"] == "mpn" and tpl["barcode_types"] == {"mpn": "ean13"}
    r = await client.post("/api/print/labels", json={"part_ids": [good["id"]], "template_id": tpl["id"]},
                          headers=H)
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    r = await client.post("/api/print/labels", json={"part_ids": [good["id"], bad["id"]],
                                                     "template_id": tpl["id"]}, headers=H)
    assert r.status_code == 422 and "Bad" in r.json()["detail"]
    # UPC-A and QR still render
    for kind in ("upca", "qr", "code128"):
        await client.put(f"/api/print/templates/{tpl['id']}", json={
            "name": "MPN", "barcode_types": {"mpn": kind}}, headers=H)
        mpn = "03600029145" if kind == "upca" else "400638133393"
        p = await _product(client, H, name=f"P-{kind}", mpn=mpn + ("" if kind == "upca" else kind))
        r = await client.post("/api/print/labels", json={"part_ids": [p["id"]], "template_id": tpl["id"]},
                              headers=H)
        assert r.status_code == 200, (kind, r.text)
    # a partial PUT kept the field
    t2 = [t for t in (await client.get("/api/print/templates", headers=H)).json()["items"] if t["id"] == tpl["id"]][0]
    assert t2["barcode_field"] == "mpn"
    # validation, permission, tenant
    assert (await client.post("/api/print/templates", json={"name": "X", "barcode_type": "pdf417"},
                              headers=H)).status_code == 422
    _, T = await _teammate(client, H, "can_manage_settings")
    assert (await client.post("/api/print/templates", json={"name": "X"}, headers=T)).status_code == 403
    H2, _, _ = await _register(client, "Other")
    assert (await client.put(f"/api/print/templates/{tpl['id']}", json={"name": "X"},
                             headers=H2)).status_code == 404


async def test_mpn_is_a_product_form_field(client):
    """MPN is a built-in on the product form (Settings → Forms), also for a saved layout, and
    it can be made required."""
    H, _, _ = await _register(client)
    items = (await client.get("/api/forms/layout?form=product", headers=H)).json()["items"]
    assert "sys:mpn" in [i["ref"] for i in items]
    # a layout saved before MPN existed still gets it
    saved = [{"ref": i["ref"], "visible": True} for i in items if i["ref"] != "sys:mpn"]
    assert (await client.put("/api/forms/layout", json={"form": "product", "items": saved}, headers=H)).status_code == 200
    items = (await client.get("/api/forms/layout?form=product", headers=H)).json()["items"]
    assert "sys:mpn" in [i["ref"] for i in items]
    # required MPN is enforced on create
    lay = [{"ref": i["ref"], "visible": True, "required": i["ref"] == "sys:mpn"} for i in items]
    assert (await client.put("/api/forms/layout", json={"form": "product", "items": lay}, headers=H)).status_code == 200
    r = await client.post("/api/catalog/products", json={"name": "No MPN"}, headers=H)
    assert r.status_code in (400, 422), r.text
    assert (await client.post("/api/catalog/products", json={"name": "Has MPN", "mpn": "M-1"},
                              headers=H)).status_code == 201


async def test_mpn_uniqueness_on_import_and_posting_commit(client):
    import openpyxl
    H, _, _ = await _register(client)
    await _product(client, H, name="Owner", mpn="UNQ-1")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "MPN"])
    ws.append(["Other", "unq-1"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await client.post("/api/catalog/products/import", headers=H,
                          files={"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 409 and "Row 2" in r.json()["detail"]
    # two rows of one sheet can't share an MPN either
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "MPN"])
    ws.append(["A1", "DUP-9"])
    ws.append(["A2", "DUP-9"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await client.post("/api/catalog/products/import", headers=H,
                          files={"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 409 and "Row 3" in r.json()["detail"]

    loc = (await client.get("/api/staff/locations", headers=H)).json()["items"][0]["id"]
    r = await client.post("/api/inventory/postings/commit", json={"location_id": loc, "rows": [
        {"new_name": "Clone", "new_mpn": "UNQ-1", "qty": 1}]}, headers=H)
    assert r.status_code == 409


async def test_scan_matches_mpn_any_case(client):
    H, _, _ = await _register(client)
    p = await _product(client, H, mpn="abc-77")
    scan = (await client.get("/api/inventory/scan?code=ABC-77", headers=H)).json()
    assert scan["product"]["id"] == p["id"] and scan["matched_on"] == "mpn"


async def test_print_template_without_field_keeps_sku_fallback(client):
    H, _, _ = await _register(client)
    tpl = (await client.post("/api/print/templates", json={"name": "Legacy"}, headers=H)).json()
    assert tpl["barcode_field"] is None
    p = await _product(client, H, sku="ONLY-SKU")
    r = await client.post("/api/print/labels", json={"part_ids": [p["id"]], "template_id": tpl["id"]}, headers=H)
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


async def test_label_preview_uses_chosen_barcode_source(client):
    H, _, _ = await _register(client)
    await _product(client, H, name="Prev", mpn="590123412345")
    r = await client.post("/api/documents/preview", json={
        "body_html": "<div>{Barcode}</div>", "doc_type": "label",
        "barcode_field": "mpn", "barcode_types": {"mpn": "ean13"}}, headers=H)
    assert r.status_code == 200 and "5901234123457" in r.json()["html"] and r.json()["sample"] == "Prev"
    r = await client.post("/api/documents/preview", json={
        "body_html": "<div>{Barcode}</div>", "doc_type": "label",
        "barcode_field": "mpn", "barcode_types": {"mpn": "upca"}}, headers=H)
    assert "be printed with this barcode type" in r.json()["html"] and "Prev" in r.json()["html"]


async def test_mpn_unique_index_backs_the_check(client, db):
    """A save that skips the pre-check (a concurrent one) is still refused by the DB index, with a 409."""
    from unittest.mock import patch

    from backend.models import Part
    H, org, _ = await _register(client)
    await _product(client, H, mpn="RACE-1")

    async def no_check(db_, org_id, mpn, exclude_id=None):
        return (mpn or "").strip() or None
    with patch("backend.routers.catalog_products._checked_mpn", no_check):
        r = await client.post("/api/catalog/products", json={"name": "Racer", "mpn": "race-1"}, headers=H)
    assert r.status_code == 409, r.text
    # archived products don't hold their MPN
    old = (await db.execute(select(Part).where(Part.org_id == org, Part.mpn == "RACE-1"))).scalar_one()
    old.active = False
    await db.commit()
    assert (await client.post("/api/catalog/products", json={"name": "Reuse", "mpn": "RACE-1"},
                              headers=H)).status_code == 201


def _sheet(*rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": ("p.xlsx", buf.getvalue(), "application/octet-stream")}


async def test_import_mpn_swap_long_and_zero(client):
    H, _, _ = await _register(client)
    a = await _product(client, H, name="Alpha", mpn="X-1")
    b = await _product(client, H, name="Beta", mpn="Y-1")
    # swapping two MPNs in one sheet is valid
    r = await client.post("/api/catalog/products/import", headers=H,
                          files=_sheet(["Name", "MPN"], ["Alpha", "Y-1"], ["Beta", "X-1"]))
    assert r.status_code == 200, r.text
    get = lambda p: client.get(f"/api/catalog/products/{p['id']}", headers=H)  # noqa: E731
    assert (await get(a)).json()["product"]["mpn"] == "Y-1"
    assert (await get(b)).json()["product"]["mpn"] == "X-1"
    # too long is refused with the row, not truncated
    r = await client.post("/api/catalog/products/import", headers=H,
                          files=_sheet(["Name", "MPN"], ["Alpha", "Z" * 70]))
    assert r.status_code == 409 and "Row 2" in r.json()["detail"]
    assert (await get(a)).json()["product"]["mpn"] == "Y-1"
    # a numeric 0 is a value, not a blank
    r = await client.post("/api/catalog/products/import", headers=H, files=_sheet(["Name", "MPN"], ["Alpha", 0]))
    assert r.status_code == 200, r.text
    assert (await get(a)).json()["product"]["mpn"] == "0"


async def test_preview_keeps_asset_labels_and_validates_options(client):
    H, _, _ = await _register(client)
    await _product(client, H, name="Some part", mpn="P-1")
    r = await client.post("/api/documents/preview", json={
        "body_html": "<p>{Asset Model}</p>", "doc_type": "vehicle_label"}, headers=H)
    assert r.status_code == 200 and r.json()["sample"] != "Some part"
    r = await client.post("/api/documents/preview", json={
        "body_html": "<p>{Asset Model}</p>", "doc_type": "label"}, headers=H)
    assert r.json()["sample"] != "Some part"
    r = await client.post("/api/documents/preview", json={
        "body_html": "<p>{Barcode}</p>", "doc_type": "label", "barcode_types": {"mpn": "ean8"}}, headers=H)
    assert r.status_code == 422
