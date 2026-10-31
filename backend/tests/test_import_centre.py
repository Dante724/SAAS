"""Import centre (inventory council items 5–8): one pipeline for Products, Services, Clients,
Assets and Suppliers — template → parse → map → preview (dry run) → commit → undo.

Covers: template round-trip, per-row errors that don't abort good rows, dry run writes nothing,
upsert by each match key, Excel quirks (leading zeros, 1.23E+12, decimal commas, BOM, date
serials), cross-shop isolation, permissions, cost columns, audit row and undo.
"""
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func, select

from backend.models import (
    AuditLog,
    Client,
    Device,
    Employee,
    ImportBatch,
    Organisation,
    Part,
    Role,
    SerializedItem,
    Service,
    StockLevel,
    Supplier,
)
from backend.services.auth import hash_password

pytestmark = pytest.mark.asyncio


# ── helpers ──────────────────────────────────────────────────────────────────
def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _parse(client, h, kind, data: bytes, name="sheet.xlsx", **params):
    r = await client.post("/api/imports/parse", headers=h, params={"kind": kind, **params},
                          files={"file": (name, data, "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


async def _run(client, h, kind, parsed, *, commit=False, mapping=None, options=None):
    body = {"headers": parsed["headers"], "rows": parsed["rows"],
            "mapping": mapping if mapping is not None else parsed["suggested"],
            "options": options or {}, "file_name": parsed.get("file_name")}
    return await client.post(f"/api/imports/{kind}/{'commit' if commit else 'preview'}", headers=h, json=body)


async def _import(client, h, kind, rows, *, name="sheet.xlsx", options=None, raw=None):
    parsed = await _parse(client, h, kind, raw if raw is not None else _xlsx(rows), name)
    r = await _run(client, h, kind, parsed, commit=True, options=options)
    assert r.status_code == 200, r.text
    return r.json()


async def _me(client, h):
    return (await client.get("/api/auth/me", headers=h)).json()


async def _count(db, model, org_id):
    return (await db.execute(select(func.count()).select_from(model).where(model.org_id == org_id))).scalar()


async def _staff_with(client, db, org_id, perms: dict):
    role = Role(org_id=org_id, name=f"R-{uuid.uuid4().hex[:5]}", permissions=perms, perms_explicit=True)
    db.add(role)
    await db.flush()
    email = f"st-{uuid.uuid4().hex[:6]}@imp.test"
    db.add(Employee(org_id=org_id, name="Staff", email=email, password_hash=hash_password("pw-staff-12345"),
                    active=True, role_id=role.id))
    await db.commit()
    r = await client.post("/api/auth/token", data={"username": email, "password": "pw-staff-12345"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


async def _second_tenant(client):
    email = f"admin2-{uuid.uuid4().hex[:8]}@test.dev"
    r = await client.post("/api/auth/register", json={
        "org_name": "Other Co", "name": "Other", "email": email, "password": "pw-test-12345"})
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


async def _set_key(db, org_id, key):
    org = (await db.execute(select(Organisation).where(Organisation.id == org_id))).scalar_one()
    org.inv_unique_identifier = key
    await db.commit()


# ── templates ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("kind", ["products", "services", "clients", "assets", "suppliers"])
async def test_template_is_a_real_workbook(client, tenant, kind):
    r = await client.get(f"/api/imports/{kind}/template", headers=tenant)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Instructions" in wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    assert any(h and h.endswith("*") for h in header), "required columns are marked"
    assert ws.max_row == 2, "exactly one example row"
    csv = await client.get(f"/api/imports/{kind}/template?fmt=csv", headers=tenant)
    assert csv.status_code == 200 and csv.content.decode("utf-8-sig").splitlines()[0].startswith(header[0])


async def test_product_template_has_dropdowns_and_text_codes(client, tenant):
    await client.post("/api/catalog/categories", headers=tenant, json={"name": "Screens", "entity_type": "product"})
    r = await client.get("/api/imports/products/template", headers=tenant)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Products"]
    header = [str(c.value).rstrip(" *") for c in ws[1]]
    for col in ("SKU", "Barcode"):
        letter = openpyxl.utils.get_column_letter(header.index(col) + 1)
        assert ws[f"{letter}2"].number_format == "@" and ws.column_dimensions[letter].number_format == "@"
    formulas = " ".join(str(dv.formula1) for dv in ws.data_validations.dataValidation)
    assert "pcs" in formulas and "yes" in formulas.lower()     # unit + yes/no lists
    assert "Lists" in formulas                                  # category + product type come from the Lists sheet
    lists = [c.value for c in wb["Lists"]["A"]] + [c.value for c in wb["Lists"]["B"]]
    assert "Screens" in lists and "Serialized" in lists


async def test_template_round_trip_products(client, tenant):
    """download → fill → import → export matches."""
    await client.post("/api/catalog/categories", headers=tenant, json={"name": "Screens", "entity_type": "product"})
    r = await client.get("/api/imports/products/template", headers=tenant)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Products"]
    header = [str(c.value).rstrip(" *") for c in ws[1]]
    ws.delete_rows(2)
    fill = [
        {"Name": "iPhone 13 Screen", "SKU": "0012-SCR", "Barcode": "0012345678905", "Category": "Screens",
         "Unit": "pcs", "Cost price": 40, "Sell price": 99.5, "Min stock": 2, "Warranty (days)": 90},
        {"Name": "Galaxy S22 Battery", "SKU": "BAT-S22", "Barcode": "", "Category": "",
         "Unit": "pcs", "Cost price": 12, "Sell price": 35, "Min stock": 0, "Warranty (days)": 30},
    ]
    for row in fill:
        ws.append([row.get(h, None) for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    out = await _import(client, tenant, "products", None, raw=buf.getvalue())
    assert out["created"] == 2 and out["failed"] == 0, out

    ex = await client.get("/api/catalog/products/export", headers=tenant)
    exp = openpyxl.load_workbook(io.BytesIO(ex.content)).active
    eh = [c.value for c in exp[1]]
    got = {r[eh.index("Name")]: dict(zip(eh, r)) for r in exp.iter_rows(min_row=2, values_only=True)}
    a = got["iPhone 13 Screen"]
    assert a["SKU"] == "0012-SCR" and a["Barcode"] == "0012345678905" and a["Category"] == "Screens"
    assert float(a["Sell price"]) == 99.5 and float(a["Cost price"]) == 40 and a["Min stock"] == 2
    assert a["Client warranty (days)"] == 90
    assert got["Galaxy S22 Battery"]["SKU"] == "BAT-S22"
    # the export re-imports onto the same products, not duplicates — unchanged rows are skipped
    again = await _import(client, tenant, "products", None, raw=ex.content, name="products.xlsx")
    assert again["created"] == 0 and again["counts"]["skip"] == 2, again
    assert all("No changes" in r["warnings"] for r in again["rows"])


async def test_example_row_left_in_the_template_is_skipped(client, tenant, db):
    r = await client.get("/api/imports/products/template", headers=tenant)
    out = await _import(client, tenant, "products", None, raw=r.content)
    assert out["created"] == 0
    me = await _me(client, tenant)
    assert await _count(db, Part, me["org_id"]) == 0


# ── per-row errors, dry run ──────────────────────────────────────────────────
async def test_bad_rows_are_reported_and_good_rows_still_import(client, tenant, db):
    rows = [["Name", "SKU", "Sell price", "Category", "Unit", "Serialized"],
            ["Good one", "G-1", "10", "", "pcs", "no"],
            ["", "NO-NAME", "5", "", "pcs", "no"],                      # missing required
            ["Bad price", "B-1", "ten", "", "pcs", "no"],              # wrong type
            ["Bad cat", "B-2", "1", "Nope", "pcs", "no"],              # unknown category
            ["Dup A", "DUP", "1", "", "pcs", "no"],
            ["Dup B", "DUP", "1", "", "pcs", "no"],                    # duplicate key in file
            ["Bad unit", "B-3", "1", "", "bucket", "no"],             # not in list
            ["Good two", "G-2", "1,50", "", "pcs", "maybe"]]          # bad yes/no
    parsed = await _parse(client, tenant, "products", _xlsx(rows))
    pre = (await _run(client, tenant, "products", parsed)).json()
    by_row = {r["row"]: r for r in pre["rows"]}
    assert by_row[2]["action"] == "create"
    assert by_row[3]["action"] == "error" and "Name" in by_row[3]["errors"][0]
    assert by_row[4]["action"] == "error" and "Sell price" in by_row[4]["errors"][0]
    assert by_row[5]["action"] == "error" and "category" in by_row[5]["errors"][0].lower()
    assert by_row[6]["action"] == "create" and by_row[7]["action"] == "error" and "row 6" in by_row[7]["errors"][0]
    assert by_row[8]["action"] == "error" and "Unit" in by_row[8]["errors"][0]
    assert by_row[9]["action"] == "error" and "yes or no" in by_row[9]["errors"][0]
    assert pre["counts"] == {"create": 2, "update": 0, "error": 6, "skip": 0}

    out = (await _run(client, tenant, "products", parsed, commit=True)).json()
    assert out["created"] == 2 and out["failed"] == 6
    me = await _me(client, tenant)
    names = set((await db.execute(select(Part.name).where(Part.org_id == me["org_id"]))).scalars())
    assert names == {"Good one", "Dup A"}


async def test_unknown_category_can_be_created_on_request(client, tenant, db):
    rows = [["Name", "Category"], ["Cable", "Accessories > Cables"]]
    parsed = await _parse(client, tenant, "products", _xlsx(rows))
    r = (await _run(client, tenant, "products", parsed, commit=True, options={"create_categories": True})).json()
    assert r["created"] == 1 and r["failed"] == 0
    tree = (await client.get("/api/catalog/categories?entity_type=product", headers=tenant)).json()
    acc = [c for c in tree["tree"] if c["name"] == "Accessories"]
    assert acc and acc[0]["children"][0]["name"] == "Cables"


async def test_dry_run_writes_nothing(client, tenant, db):
    me = await _me(client, tenant)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    rows = [["Name", "SKU", "Opening stock", "Serial numbers", "Serialized", "Category"],
            ["Phone", "PH-1", "2", "SN1, SN2", "yes", "New cat"]]
    parsed = await _parse(client, tenant, "products", _xlsx(rows))
    for _ in range(2):
        r = await _run(client, tenant, "products", parsed,
                       options={"location_id": loc, "create_categories": True})
        assert r.status_code == 200 and r.json()["counts"]["create"] == 1
    assert await _count(db, Part, me["org_id"]) == 0
    assert await _count(db, SerializedItem, me["org_id"]) == 0
    assert await _count(db, ImportBatch, me["org_id"]) == 0
    cats = (await client.get("/api/catalog/categories?entity_type=product", headers=tenant)).json()
    assert "New cat" not in str(cats)
    for kind, rws in [("clients", [["Name", "Phone"], ["Ann", "0123"]]),
                      ("services", [["Name", "Price"], ["Diag", "10"]]),
                      ("suppliers", [["Name"], ["Acme"]]),
                      ("assets", [["Brand", "Model", "Serial"], ["Apple", "iPhone", "X1"]])]:
        p = await _parse(client, tenant, kind, _xlsx(rws))
        assert (await _run(client, tenant, kind, p)).status_code == 200
    assert await _count(db, Client, me["org_id"]) == 0
    assert await _count(db, Service, me["org_id"]) == 0
    assert await _count(db, Supplier, me["org_id"]) == 0
    assert await _count(db, Device, me["org_id"]) == 0


# ── upsert by each key ───────────────────────────────────────────────────────
async def _seed_part(db, org_id, **kw):
    p = Part(org_id=org_id, **kw)
    db.add(p)
    await db.commit()
    return p


@pytest.mark.parametrize("key,row,expect_update", [
    ("sku", ["Renamed", "SKU-1", "999", "50"], True),           # matched on SKU though renamed
    ("sku", ["Original", "SKU-X", "999", "50"], False),          # name alone doesn't match under SKU
    ("barcode", ["Renamed", "SKU-X", "4006381333931", "50"], True),
    ("barcode", ["Renamed", "SKU-1", "999", "50"], False),
    ("name", ["original", "SKU-X", "999", "50"], True),          # case-insensitive name
    ("code", ["Renamed", "SKU-X", "999", "50"], False),
])
async def test_upsert_follows_the_unique_identifier_setting(client, tenant, db, key, row, expect_update):
    me = await _me(client, tenant)
    await _seed_part(db, me["org_id"], name="Original", sku="SKU-1", barcode="4006381333931", code="C-1")
    await _set_key(db, me["org_id"], key)
    out = await _import(client, tenant, "products", [["Name", "SKU", "Barcode", "Sell price"], row])
    assert out["updated" if expect_update else "created"] == 1, out


async def test_default_key_is_sku_then_barcode_then_name(client, tenant, db):
    me = await _me(client, tenant)
    await _seed_part(db, me["org_id"], name="A", sku="S-A")
    await _seed_part(db, me["org_id"], name="B", barcode="4006381333931")
    await _seed_part(db, me["org_id"], name="C")
    parsed = await _parse(client, tenant, "products", _xlsx([
        ["Name", "SKU", "Barcode", "Sell price"],
        ["A renamed", "s-a", "", "1"],                # SKU, case-insensitive
        ["B renamed", "", "4006381333931", "2"],      # barcode
        ["c", "", "", "3"],                           # name
        ["D", "S-D", "", "4"]]))
    pre = (await _run(client, tenant, "products", parsed)).json()
    assert [r["action"] for r in pre["rows"]] == ["update", "update", "update", "create"]
    assert [r["matched_on"] for r in pre["rows"][:3]] == ["sku", "barcode", "name"]


async def test_clients_match_phone_then_email_and_services_by_name(client, tenant, db):
    me = await _me(client, tenant)
    db.add(Client(org_id=me["org_id"], name="Ann", phone="+1 (555) 010-0001"))
    db.add(Client(org_id=me["org_id"], name="Bob", email="bob@x.test"))
    db.add(Service(org_id=me["org_id"], name="Screen fit", price=10))
    await db.commit()
    out = await _import(client, tenant, "clients", [["Name", "Phone", "Email"],
                                                    ["Ann B", "15550100001", ""],
                                                    ["Robert", "", "BOB@x.test"],
                                                    ["Cara", "555-0200", ""]])
    assert (out["created"], out["updated"]) == (1, 2), out
    out = await _import(client, tenant, "services", [["Name", "Price", "Type"], ["screen FIT", "25", "Labor"]])
    assert out["updated"] == 1
    s = (await db.execute(select(Service).where(Service.org_id == me["org_id"]))).scalar_one()
    await db.refresh(s)
    assert float(s.price) == 25 and s.is_labor is True


# ── Excel quirks ─────────────────────────────────────────────────────────────
async def test_excel_quirks_leading_zeros_scientific_decimal_comma_bom_dates(client, tenant, db):
    me = await _me(client, tenant)
    # BOM + semicolon CSV with decimal commas, the shape European Excel saves
    csv = "﻿Name;SKU;Barcode;Sell price;Cost price\r\n" \
          "Lens;00123;012345678905;12,50;1.234,56\r\n" \
          "Sci;S-2;1.23457E+12;3;1\r\n" \
          "Short UPC;S-3;36000291452;3;1\r\n"
    parsed = await _parse(client, tenant, "products", csv.encode("utf-8"), name="p.csv")
    assert parsed["headers"][0] == "Name"
    pre = (await _run(client, tenant, "products", parsed)).json()
    rows = {r["row"]: r for r in pre["rows"]}
    assert rows[3]["action"] == "error" and "E+" in rows[3]["errors"][0]
    out = (await _run(client, tenant, "products", parsed, commit=True)).json()
    assert out["created"] == 2
    lens = (await db.execute(select(Part).where(Part.org_id == me["org_id"], Part.name == "Lens"))).scalar_one()
    assert lens.sku == "00123" and lens.barcode == "012345678905"
    assert float(lens.sell_price) == 12.5 and float(lens.cost_price) == 1234.56
    upc = (await db.execute(select(Part).where(Part.org_id == me["org_id"], Part.name == "Short UPC"))).scalar_one()
    assert upc.barcode == "036000291452", "an 11-digit UPC that lost its leading zero is re-padded"

    # XLSX: a barcode stored as a number (Excel showed 1.23457E+12) keeps every digit
    raw = _xlsx([["Name", "Barcode"], ["Num code", 4006381333931]])
    parsed = await _parse(client, tenant, "products", raw)
    assert parsed["rows"][0][1] == "4006381333931"


async def test_numeric_barcode_matches_despite_lost_leading_zero(client, tenant, db):
    me = await _me(client, tenant)
    await _set_key(db, me["org_id"], "barcode")
    await _seed_part(db, me["org_id"], name="Lens", barcode="0012345678905")
    out = await _import(client, tenant, "products", [["Name", "Barcode", "Sell price"], ["Lens", 12345678905, 5]])
    assert out["updated"] == 1, out


async def test_dates_from_excel_serials_and_text(client, tenant, db):
    import datetime as dt
    me = await _me(client, tenant)
    db.add(Client(org_id=me["org_id"], name="Owner", phone="5550001"))
    await db.commit()
    raw = _xlsx([["Brand", "Model", "Serial", "Owner phone", "Purchase date"],
                 ["Apple", "iPhone 12", "A1", "5550001", 45292],             # serial number → 2024-01-01
                 ["Apple", "iPhone 13", "A2", "5550001", dt.datetime(2024, 2, 3)],
                 ["Apple", "iPhone 14", "A3", "", "15.03.2024"],
                 ["Apple", "iPhone 15", "A4", "", "not a date"]])
    parsed = await _parse(client, tenant, "assets", raw)
    out = (await _run(client, tenant, "assets", parsed, commit=True)).json()
    assert out["created"] == 3 and out["failed"] == 1, out
    got = {d.serial: d for d in (await db.execute(select(Device).where(Device.org_id == me["org_id"]))).scalars()}
    assert got["A1"].purchase_date == dt.date(2024, 1, 1)
    assert got["A2"].purchase_date == dt.date(2024, 2, 3)
    assert got["A3"].purchase_date == dt.date(2024, 3, 15) and got["A3"].owner_type == "company"
    assert got["A1"].client_id is not None and got["A1"].owner_type == "client"


# ── opening stock, bins, serials ─────────────────────────────────────────────
async def test_opening_stock_bin_and_serials(client, tenant, db):
    me = await _me(client, tenant)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    b = await client.post("/api/bins", headers=tenant, json={"location_id": loc, "name": "A1"})
    assert b.status_code in (200, 201), b.text
    rows = [["Name", "SKU", "Serialized", "Opening stock", "Bin", "Serial numbers", "Cost price"],
            ["Phone X", "PX", "yes", "", "A1", "IMEI-1; IMEI-2", "100"],      # qty = number of serials
            ["Cable", "CB", "no", "10", "A1", "", "2"],
            ["Phone Y", "PY", "yes", "2", "", "Y-1", ""],                        # 1 serial for 2 units
            ["Phone Z", "PZ", "yes", "1", "Nowhere", "", ""]]                    # unknown bin
    parsed = await _parse(client, tenant, "products", _xlsx(rows))
    pre = (await _run(client, tenant, "products", parsed, options={"location_id": loc})).json()
    by = {r["row"]: r for r in pre["rows"]}
    assert by[4]["action"] == "error" and "serial" in by[4]["errors"][0].lower()
    assert by[5]["action"] == "error" and "bin" in by[5]["errors"][0].lower()
    # generate the missing serial instead
    out = (await _run(client, tenant, "products", parsed, commit=True,
                      options={"location_id": loc, "generate_serials": True})).json()
    assert out["created"] == 3 and out["failed"] == 1, out
    px = (await db.execute(select(Part).where(Part.org_id == me["org_id"], Part.sku == "PX"))).scalar_one()
    sl = (await db.execute(select(StockLevel).where(StockLevel.part_id == px.id))).scalar_one()
    assert sl.qty_available == 2 and sl.bin_id is not None
    sns = set((await db.execute(select(SerializedItem.serial_number).where(SerializedItem.part_id == px.id))).scalars())
    assert sns == {"IMEI-1", "IMEI-2"}
    py = (await db.execute(select(Part).where(Part.org_id == me["org_id"], Part.sku == "PY"))).scalar_one()
    py_sns = list((await db.execute(select(SerializedItem.serial_number).where(SerializedItem.part_id == py.id))).scalars())
    assert len(py_sns) == 2 and "Y-1" in py_sns and len(set(py_sns)) == 2
    cb = (await db.execute(select(Part).where(Part.org_id == me["org_id"], Part.sku == "CB"))).scalar_one()
    assert (await db.execute(select(StockLevel.qty_available).where(StockLevel.part_id == cb.id))).scalar() == 10


async def test_serial_already_in_stock_is_a_row_error(client, tenant, db):
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    await _import(client, tenant, "products", [["Name", "SKU", "Serialized", "Serial numbers"], ["P", "P1", "yes", "S-1"]],
                  options={"location_id": loc})
    parsed = await _parse(client, tenant, "products", _xlsx([
        ["Name", "SKU", "Serialized", "Serial numbers"], ["Q", "Q1", "yes", "S-1"], ["R", "R1", "yes", "S-9, S-9"]]))
    pre = (await _run(client, tenant, "products", parsed, options={"location_id": loc})).json()
    assert all(r["action"] == "error" for r in pre["rows"])
    assert "already in stock" in pre["rows"][0]["errors"][0] and "twice" in pre["rows"][1]["errors"][0]


# ── cross-shop + permissions ─────────────────────────────────────────────────
async def test_other_shops_products_are_invisible(client, tenant, db):
    other = await _second_tenant(client)
    await _import(client, other, "products", [["Name", "SKU", "Barcode"], ["Theirs", "SHARED", "4006381333931"]])
    out = await _import(client, tenant, "products", [["Name", "SKU", "Barcode"], ["Mine", "SHARED", "4006381333931"]])
    assert out["created"] == 1 and out["updated"] == 0 and out["failed"] == 0
    # categories / bins / clients of another shop aren't usable either
    await client.post("/api/catalog/categories", headers=other, json={"name": "TheirCat", "entity_type": "product"})
    parsed = await _parse(client, tenant, "products", _xlsx([["Name", "Category"], ["X", "TheirCat"]]))
    pre = (await _run(client, tenant, "products", parsed)).json()
    assert pre["rows"][0]["action"] == "error"
    oth_me = await _me(client, other)
    db.add(Client(org_id=oth_me["org_id"], name="Their client", phone="5559999"))
    await db.commit()
    parsed = await _parse(client, tenant, "assets", _xlsx([["Brand", "Serial", "Owner phone"], ["A", "Z1", "5559999"]]))
    pre = (await _run(client, tenant, "assets", parsed)).json()
    assert pre["rows"][0]["action"] == "error" and "client" in pre["rows"][0]["errors"][0].lower()


async def test_undo_is_per_shop(client, tenant):
    other = await _second_tenant(client)
    out = await _import(client, tenant, "suppliers", [["Name"], ["Acme"]])
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=other)
    assert r.status_code == 404


async def test_create_needs_can_create_products_and_update_needs_can_edit(client, tenant, db):
    me = await _me(client, tenant)
    await _seed_part(db, me["org_id"], name="Existing", sku="E-1", sell_price=5)
    rows = [["Name", "SKU", "Sell price"], ["Existing", "E-1", "9"], ["Brand new", "N-1", "3"]]
    edit_only = await _staff_with(client, db, me["org_id"], {"can_edit_products": True})
    pre = (await _run(client, edit_only, "products", await _parse(client, edit_only, "products", _xlsx(rows)))).json()
    assert [r["action"] for r in pre["rows"]] == ["update", "error"]
    assert "create products" in pre["rows"][1]["errors"][0]
    create_only = await _staff_with(client, db, me["org_id"], {"can_create_products": True})
    pre = (await _run(client, create_only, "products", await _parse(client, create_only, "products", _xlsx(rows)))).json()
    assert [r["action"] for r in pre["rows"]] == ["error", "create"]
    nobody = await _staff_with(client, db, me["org_id"], {})
    r = await client.post("/api/imports/parse?kind=products", headers=nobody,
                          files={"file": ("x.xlsx", _xlsx(rows), "application/octet-stream")})
    assert r.status_code == 403
    r = await client.post("/api/imports/products/commit", headers=nobody,
                          json={"headers": rows[0], "rows": rows[1:], "mapping": {}})
    assert r.status_code == 403


async def test_cost_columns_are_dropped_for_roles_that_cannot_see_cost(client, tenant, db):
    me = await _me(client, tenant)
    await _seed_part(db, me["org_id"], name="Existing", sku="E-1", cost_price=7, sell_price=5)
    h = await _staff_with(client, db, me["org_id"], {"can_edit_products": True, "can_create_products": True})
    parsed = await _parse(client, h, "products", _xlsx([["Name", "SKU", "Cost price", "Sell price"],
                                                         ["Existing", "E-1", "1", "9"], ["New", "N-1", "50", "60"]]))
    pre = (await _run(client, h, "products", parsed)).json()
    assert any("cost" in w.lower() for w in pre["warnings"])
    await _run(client, h, "products", parsed, commit=True)
    parts = {p.sku: p for p in (await db.execute(select(Part).where(Part.org_id == me["org_id"]))).scalars()}
    for p in parts.values():
        await db.refresh(p)
    assert float(parts["E-1"].cost_price) == 7 and float(parts["E-1"].sell_price) == 9
    assert float(parts["N-1"].cost_price) == 0


async def test_opening_stock_needs_manage_inventory(client, tenant, db):
    me = await _me(client, tenant)
    h = await _staff_with(client, db, me["org_id"], {"can_create_products": True})
    parsed = await _parse(client, h, "products", _xlsx([["Name", "Opening stock"], ["P", "5"]]))
    pre = (await _run(client, h, "products", parsed)).json()
    assert pre["rows"][0]["action"] == "error" and "stock" in pre["rows"][0]["errors"][0].lower()


async def test_other_lists_check_their_own_permissions(client, tenant, db):
    me = await _me(client, tenant)
    h = await _staff_with(client, db, me["org_id"], {"can_create_products": True})
    for kind, rows in [("clients", [["Name", "Phone"], ["A", "1"]]), ("services", [["Name"], ["S"]]),
                       ("suppliers", [["Name"], ["Acme"]]), ("assets", [["Brand", "Serial"], ["A", "1"]])]:
        r = await client.post(f"/api/imports/{kind}/preview", headers=h,
                              json={"headers": rows[0], "rows": rows[1:], "mapping": {}})
        assert r.status_code == 403, kind


# ── audit, undo, mapping memory, caps ────────────────────────────────────────
async def test_audit_row_and_undo_last_import(client, tenant, db):
    me = await _me(client, tenant)
    await _seed_part(db, me["org_id"], name="Old", sku="O-1", sell_price=5)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    out = await _import(client, tenant, "products",
                        [["Name", "SKU", "Sell price", "Opening stock", "Serialized", "Serial numbers"],
                         ["Old", "O-1", "8", "3", "no", ""], ["New", "N-1", "2", "", "yes", "U1, U2"]],
                        name="mine.xlsx", options={"location_id": loc})
    audit = (await db.execute(select(AuditLog).where(AuditLog.org_id == me["org_id"],
                                                     AuditLog.action == "import"))).scalar_one()
    assert audit.entity_type == "products" and audit.entity_id == out["batch_id"]
    assert audit.diff["file_name"] == "mine.xlsx" and audit.diff["created"] == 1 and audit.diff["updated"] == 1

    last = (await client.get("/api/imports/batches?kind=products", headers=tenant)).json()["items"][0]
    assert last["id"] == out["batch_id"] and last["can_undo"] is True
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 200, r.text
    parts = (await db.execute(select(Part).where(Part.org_id == me["org_id"]))).scalars().all()
    assert [p.name for p in parts] == ["Old"]
    await db.refresh(parts[0])
    assert float(parts[0].sell_price) == 5
    assert (await db.execute(select(func.coalesce(func.sum(StockLevel.qty_available), 0))
                             .where(StockLevel.part_id == parts[0].id))).scalar() == 0
    assert await _count(db, SerializedItem, me["org_id"]) == 0
    assert (await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)).status_code == 409


async def test_undo_refuses_when_imported_records_were_used(client, tenant, db):
    me = await _me(client, tenant)
    out = await _import(client, tenant, "clients", [["Name", "Phone"], ["Ann", "5551"]])
    c = (await db.execute(select(Client).where(Client.org_id == me["org_id"]))).scalar_one()
    db.add(Device(org_id=me["org_id"], client_id=c.id, brand="Apple"))
    await db.commit()
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 409 and "used" in r.json()["detail"]


async def test_only_the_last_import_can_be_undone(client, tenant):
    first = await _import(client, tenant, "suppliers", [["Name"], ["Acme"]])
    await _import(client, tenant, "suppliers", [["Name"], ["Beta"]])
    r = await client.post(f"/api/imports/batches/{first['batch_id']}/undo", headers=tenant)
    assert r.status_code == 409


async def test_mapping_is_remembered_per_list_and_supplier(client, tenant):
    sup = (await client.post("/api/suppliers", headers=tenant, json={"name": "Acme"})).json()
    r = await client.put(f"/api/imports/posting/mapping?supplier_id={sup['id']}", headers=tenant,
                         json={"mapping": {"Part No": "code", "Desc": "name"}})
    assert r.status_code == 200
    got = (await client.get(f"/api/imports/posting/mapping?supplier_id={sup['id']}", headers=tenant)).json()
    assert got["mapping"] == {"Part No": "code", "Desc": "name"}
    assert (await client.get("/api/imports/posting/mapping", headers=tenant)).json()["mapping"] is None
    # a list's remembered mapping is used as the suggestion on the next parse
    await client.put("/api/imports/products/mapping", headers=tenant, json={"mapping": {"Artikel": "name"}})
    parsed = await _parse(client, tenant, "products", _xlsx([["Artikel"], ["X"]]))
    assert parsed["suggested"] == {"Artikel": "name"}
    other = await _second_tenant(client)
    r = await client.get(f"/api/imports/posting/mapping?supplier_id={sup['id']}", headers=other)
    assert r.status_code == 404


async def test_auto_mapping_understands_common_headers(client, tenant):
    parsed = await _parse(client, tenant, "products", _xlsx([
        ["Product name", "Item #", "UPC", "Retail price", "Qty", "IMEIs", "Unknown"], ["a", "b", "1", "2", "3", "", ""]]))
    assert parsed["suggested"] == {"Product name": "name", "Item #": "sku", "UPC": "barcode",
                                   "Retail price": "sell_price", "Qty": "opening_qty", "IMEIs": "serials"}


async def test_caps_and_formats(client, tenant):
    from backend.core import spreadsheet
    r = await client.post("/api/imports/parse?kind=products", headers=tenant,
                          files={"file": ("x.xlsm", b"PK\x03\x04junk", "application/octet-stream")})
    assert r.status_code == 422
    big = "Name\n" + "\n".join(f"P{i}" for i in range(spreadsheet.MAX_IMPORT_ROWS + 1))
    r = await client.post("/api/imports/parse?kind=products", headers=tenant,
                          files={"file": ("x.csv", big.encode(), "text/csv")})
    assert r.status_code == 413
    r = await client.post("/api/imports/products/preview", headers=tenant, json={
        "headers": ["Name"], "rows": [["x"]] * (spreadsheet.MAX_IMPORT_ROWS + 1), "mapping": {"Name": "name"}})
    assert r.status_code == 422
    r = await client.post("/api/imports/parse?kind=products", headers=tenant,
                          files={"file": ("x.csv", b"\x00" * (spreadsheet.MAX_XLSX_BYTES + 10), "text/csv")})
    assert r.status_code == 413


async def test_formula_text_is_unguarded_on_import(client, tenant, db):
    me = await _me(client, tenant)
    await _import(client, tenant, "products", [["Name"], ["'=cmd"]])
    p = (await db.execute(select(Part).where(Part.org_id == me["org_id"]))).scalar_one()
    assert p.name == "=cmd"
    ex = await client.get("/api/catalog/products/export", headers=tenant)
    ws = openpyxl.load_workbook(io.BytesIO(ex.content)).active
    names = [r[3] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert names == ["'=cmd"], "export keeps the apostrophe guard"


async def test_suppliers_import_creates_contact_mirror(client, tenant, db):
    me = await _me(client, tenant)
    out = await _import(client, tenant, "suppliers", [["Name", "Email", "Lead time (days)"], ["Acme", "a@acme.test", "5"]])
    assert out["created"] == 1
    s = (await db.execute(select(Supplier).where(Supplier.org_id == me["org_id"]))).scalar_one()
    assert s.lead_days == 5 and s.client_id
    c = (await db.execute(select(Client).where(Client.id == s.client_id))).scalar_one()
    assert c.is_supplier
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 200
    assert await _count(db, Supplier, me["org_id"]) == 0 and await _count(db, Client, me["org_id"]) == 0


# ── postings: serials column + per-supplier mapping ──────────────────────────
async def test_posting_commit_takes_serials_and_generates_missing(client, tenant, db):
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    me = await _me(client, tenant)
    p = await _seed_part(db, me["org_id"], name="Phone", sku="PH", serialized=True)
    r = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "rows": [{"part_id": p.id, "qty": 3, "serials": ["A-1", "A-2"]}],
        "generate_serials": True})
    assert r.status_code == 201, r.text
    sns = list((await db.execute(select(SerializedItem.serial_number).where(SerializedItem.part_id == p.id))).scalars())
    assert len(sns) == 3 and {"A-1", "A-2"} <= set(sns)
    r = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "rows": [{"part_id": p.id, "qty": 1, "serials": ["A-1"]}]})
    assert r.status_code == 409 and "A-1" in r.text
    r = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "rows": [{"part_id": p.id, "qty": 2, "serials": ["B-1"]}]})
    assert r.status_code == 422 and "serial" in r.text.lower()


async def test_parse_works_for_posting_files(client, tenant):
    raw = _xlsx([["Part No", "Description", "Qty", "Unit cost", "Serials"], ["SCR-1", "Screen", 2, 10.5, "X1, X2"]])
    parsed = await _parse(client, tenant, "posting", raw)
    assert parsed["headers"] == ["Part No", "Description", "Qty", "Unit cost", "Serials"]
    assert parsed["rows"] == [["SCR-1", "Screen", "2", "10.5", "X1, X2"]]
    assert parsed["suggested"]["Serials"] == "serials"


# ── review fixes (security) ──────────────────────────────────────────────────
async def test_too_many_columns_is_refused(client, tenant):
    wide = ("," * 5000 + "x\n" + "a\n").encode()
    r = await client.post("/api/imports/parse", headers=tenant, params={"kind": "products"},
                          files={"file": ("w.csv", wide, "text/csv")})
    assert r.status_code == 413 and "columns" in r.text
    ok = await _parse(client, tenant, "products", b"Name,,,,\nA,,,,\n", name="t.csv")
    assert ok["headers"] == ["Name"], "trailing empty columns are dropped"


async def test_serial_generation_is_capped_and_template_checked(client, tenant, db):
    me = await _me(client, tenant)
    org = (await db.execute(select(Organisation).where(Organisation.id == me["org_id"]))).scalar_one()
    org.inv_product_serial_template = "SN-{C 6}"          # no real counter → falls back to SN-{C:6}
    await db.commit()
    parsed = await _parse(client, tenant, "products", _xlsx([["Name", "Serialized", "Opening stock"],
                                                             ["Big", "yes", "20000"], ["Ok", "yes", "2"]]))
    pre = (await _run(client, tenant, "products", parsed, options={"generate_serials": True})).json()
    assert pre["rows"][0]["action"] == "error" and "10,000" in pre["rows"][0]["errors"][0]
    out = (await _run(client, tenant, "products", parsed, commit=True, options={"generate_serials": True})).json()
    assert out["created"] == 1
    sns = set((await db.execute(select(SerializedItem.serial_number).where(SerializedItem.org_id == me["org_id"]))).scalars())
    assert len(sns) == 2 and all(s.startswith("SN-") for s in sns)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    p = await _seed_part(db, me["org_id"], name="Phone", serialized=True)
    r = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "rows": [{"part_id": p.id, "qty": 1e9}], "generate_serials": True})
    assert r.status_code == 422


async def test_asset_updates_need_manage_inventory(client, tenant, db):
    me = await _me(client, tenant)
    db.add(Device(org_id=me["org_id"], brand="Apple", serial="S1", imei="490154203237518"))
    await db.commit()
    h = await _staff_with(client, db, me["org_id"], {"can_view_asset_info": True})
    parsed = await _parse(client, h, "assets", _xlsx([["Brand", "IMEI", "Serial"], ["Apple", "490154203237518", "HACKED"],
                                                       ["Nokia", "", "N1"]]))
    pre = (await _run(client, h, "assets", parsed)).json()
    assert [r["action"] for r in pre["rows"]] == ["error", "create"]
    assert "edit existing assets" in pre["rows"][0]["errors"][0]


async def test_undo_keeps_values_changed_after_the_import(client, tenant, db):
    me = await _me(client, tenant)
    p = await _seed_part(db, me["org_id"], name="Lens", sku="L-1", sell_price=10, min_stock=1)
    out = await _import(client, tenant, "products", [["SKU", "Sell price", "Min stock"], ["L-1", "12", "3"]])
    r = await client.put(f"/api/catalog/products/{p.id}", headers=tenant, json={"name": "Lens", "sell_price": 15})
    assert r.status_code == 200, r.text
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 200 and r.json()["kept_newer"] == 1
    await db.refresh(p)
    assert float(p.sell_price) == 15 and p.min_stock == 1


async def test_undo_refuses_after_other_stock_moves_on_an_updated_part(client, tenant, db):
    me = await _me(client, tenant)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    p = await _seed_part(db, me["org_id"], name="Cable", sku="C-1")
    out = await _import(client, tenant, "products", [["SKU", "Opening stock", "Cost price"], ["C-1", "5", "2"]],
                        options={"location_id": loc})
    r = await client.post("/api/inventory/postings", headers=tenant, json={
        "location_id": loc, "lines": [{"part_id": p.id, "qty": 3, "purchase_price": 4}]})
    assert r.status_code == 201, r.text
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 409


async def test_template_lists_are_formula_safe(client, tenant):
    await client.post("/api/catalog/categories", headers=tenant, json={"name": '=HYPERLINK("http://x","y")', "entity_type": "product"})
    r = await client.get("/api/imports/products/template", headers=tenant)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    vals = [c.value for c in wb["Lists"]["A"]]
    assert "'=HYPERLINK(\"http://x\",\"y\")" in vals and not any(str(v).startswith("=") for v in vals if v)


# ── review fixes (tester) ────────────────────────────────────────────────────
async def test_cell_parsing_edge_cases():
    from backend.services.import_cells import CellError, to_code, to_date, to_number
    assert [to_number(v, "x") for v in ("0,125", "1,2345", "1,234", "1.000,5", "1.234.567", "-1,5", "12,50")] == \
        [0.125, 1.2345, 1234, 1000.5, 1234567, -1.5, 12.5]
    assert to_code("100E5", "SKU") == "100E5"
    for bad in ("1.23457E+12", "4E+12"):
        with pytest.raises(CellError):
            to_code(bad, "Barcode")
    import datetime as dt
    assert to_date("45292.5", "d") == dt.date(2024, 1, 1)
    assert to_date("2024-01-02T10:00:00Z", "d") == dt.date(2024, 1, 2)
    with pytest.raises(CellError):
        to_date("2024", "d")


async def test_undo_removes_categories_the_import_created(client, tenant, db):
    from backend.models import Category
    me = await _me(client, tenant)
    p = await _seed_part(db, me["org_id"], name="Old", sku="O-1")
    out = await _import(client, tenant, "products", [["Name", "SKU", "Category"], ["X", "X-1", "New > Sub"],
                                                     ["Old", "O-1", "New > Sub"]], options={"create_categories": True})
    await db.refresh(p)
    assert p.category == "Sub"
    r = await client.post(f"/api/imports/batches/{out['batch_id']}/undo", headers=tenant)
    assert r.status_code == 200, r.text
    await db.refresh(p)
    assert p.category_id is None and p.category is None
    assert await _count(db, Category, me["org_id"]) == 0


async def test_example_row_is_skipped_after_excel_strips_zeros(client, tenant, db):
    r = await client.get("/api/imports/products/template?fmt=csv", headers=tenant)
    text = r.content.decode("utf-8-sig").replace("0012345678905", "12345678905")
    out = await _import(client, tenant, "products", None, raw=text.encode(), name="t.csv")
    assert out["created"] == 0 and out["counts"]["skip"] == 1


async def test_ignored_columns_are_remembered(client, tenant):
    parsed = await _parse(client, tenant, "products", _xlsx([["Name", "SKU"], ["A", "S-1"]]))
    r = await _run(client, tenant, "products", parsed, commit=True, mapping={"Name": "name"})
    assert r.status_code == 200
    again = await _parse(client, tenant, "products", _xlsx([["Name", "SKU"], ["B", "S-2"]]))
    assert again["suggested"] == {"Name": "name"}


async def test_posting_generated_serials_never_repeat_typed_ones(client, tenant, db):
    me = await _me(client, tenant)
    loc = (await client.get("/api/staff/locations", headers=tenant)).json()["items"][0]["id"]
    p = await _seed_part(db, me["org_id"], name="Phone", serialized=True)
    first = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "rows": [{"part_id": p.id, "qty": 1}], "generate_serials": True})
    assert first.status_code == 201
    gen = (await db.execute(select(SerializedItem.serial_number).where(SerializedItem.part_id == p.id))).scalar_one()
    nxt = gen[:-1] + str(int(gen[-1]) + 1)        # what the generator would make next
    r = await client.post("/api/inventory/postings/commit", headers=tenant, json={
        "location_id": loc, "generate_serials": True,
        "rows": [{"part_id": p.id, "qty": 1}, {"part_id": p.id, "qty": 1, "serials": [nxt]}]})
    assert r.status_code == 201, r.text
    sns = list((await db.execute(select(SerializedItem.serial_number).where(SerializedItem.part_id == p.id))).scalars())
    assert len(sns) == 3 and len(set(sns)) == 3
