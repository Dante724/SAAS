"""Import centre — one pipeline for every list (inventory council items 5–8).

    parse file → map columns → preview (dry run, per-row verdicts) → commit → undo

Each list (Products, Services, Clients, Assets, Suppliers) is a `Kind`: its columns (`F`), who may
use it, how a row finds the record it updates, and how a row is saved. The runner does the rest the
same way for all of them:

- every row gets its own verdict — create, update, skip or error — and an error never stops the
  good rows (commit saves each row in a SAVEPOINT);
- preview runs the exact same checks as commit and writes nothing;
- cost columns are dropped for roles that can't see costs; creating and editing need the list's
  own permissions;
- a commit writes one ImportBatch (what the undo needs) and one audit row.

Products match on the shop's "unique identifier" setting (organisations.inv_unique_identifier):
SKU, barcode, code or name. Unset means SKU, then barcode, then name.
"""
import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any, Optional

from fastapi import HTTPException
from sqlalchemy import Date, and_, delete, func, not_, or_, select, update

from backend.core.spreadsheet import safe_cell, unguard_cell
from backend.models import (
    Category,
    Client,
    Device,
    Employee,
    ImportBatch,
    Location,
    Organisation,
    Part,
    SerializedItem,
    Service,
    StockItemBatch,
    StockLevel,
    StockMovement,
    StorageBin,
    Supplier,
)
from backend.services import import_cells as C
from backend.services.auth import has_perm
from backend.services.money import q2

UNDO_WINDOW = timedelta(hours=24)
COST_PERM = "can_view_purchase_price"
MAX_OPENING_QTY = 100_000
MAX_SERIALS = 10_000          # serial numbers (given or generated) in one import or posting


# ── columns ──────────────────────────────────────────────────────────────────
@dataclass(frozen=True)
class F:
    key: str
    label: str
    type: str = "text"          # text|code|barcode|int|num|bool|date|choice|serials
    required: bool = False      # required to create a record
    choices: tuple = ()
    synonyms: tuple = ()
    help: str = ""
    example: Any = ""
    cost: bool = False          # dropped for roles that can't see costs
    maxlen: Optional[int] = None
    in_template: bool = True
    list_name: Optional[str] = None   # dropdown from the Lists sheet (categories)


def parse_cell(f: F, raw: str):
    if f.type == "num":
        v = C.to_number(raw, f.label)
        if v < 0:
            raise C.CellError(f"{f.label} can't be negative")
        return v
    if f.type == "int":
        v = C.to_int(raw, f.label)
        if v < 0:
            raise C.CellError(f"{f.label} can't be negative")
        return v
    if f.type == "bool":
        return C.to_bool(raw, f.label)
    if f.type == "date":
        return C.to_date(raw, f.label)
    if f.type in ("code", "barcode"):
        v = C.to_code(raw, f.label, barcode=f.type == "barcode")
    elif f.type == "serials":
        return C.to_serials(raw)
    elif f.type == "choice":
        for c in f.choices:
            if raw.strip().lower() == c.lower():
                return c
        raise C.CellError(f"{f.label} must be one of: {', '.join(f.choices)} (got '{raw}')")
    else:
        v = raw.strip()
    if f.maxlen and len(v) > f.maxlen:
        raise C.CellError(f"{f.label} is longer than {f.maxlen} characters")
    return v


# ── one row's verdict ────────────────────────────────────────────────────────
@dataclass
class Row:
    row: int
    rec: dict
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    action: str = "create"      # create|update|skip|error
    target: Any = None
    matched_on: Optional[str] = None
    label: str = ""
    changes: dict = field(default_factory=dict)
    extra: dict = field(default_factory=dict)

    def out(self) -> dict:
        return {"row": self.row, "action": self.action, "label": self.label, "matched_on": self.matched_on,
                "errors": self.errors, "warnings": self.warnings, "changes": sorted(self.changes)}


@dataclass
class Ctx:
    db: Any
    current: Employee
    org: Organisation
    options: dict
    perms: dict
    seen: dict = field(default_factory=dict)        # (what, key) -> sheet row
    batch_id: Optional[str] = None
    undo: dict = field(default_factory=lambda: {"created": [], "updated": [], "stock": []})
    data: dict = field(default_factory=dict)

    def claim(self, what: str, key: str, row: Row, shown: str) -> bool:
        k = (what, key)
        if k in self.seen and self.seen[k] != row.row:
            row.errors.append(f"{shown} is also on row {self.seen[k]}")
            return False
        self.seen[k] = row.row
        return True


def _jsonable(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return v


def _same(a, b) -> bool:
    if isinstance(a, (int, float, Decimal)) and isinstance(b, (int, float, Decimal)) and not isinstance(a, bool):
        return abs(float(a) - float(b)) < 1e-9
    return (a if a != "" else None) == (b if b != "" else None)


def _changes(target, values: dict) -> dict:
    return {k: v for k, v in values.items() if not _same(getattr(target, k), v)}


def _apply_update(ctx: Ctx, table: str, target, changes: dict, extra_before: Optional[dict] = None):
    """Write `changes` and remember the old values (and the new ones, so undo only restores a field
    nobody has touched since)."""
    before = {k: _jsonable(getattr(target, k)) for k in changes}
    before.update(extra_before or {})
    for k, v in changes.items():
        setattr(target, k, v)
    if before:
        ctx.undo["updated"].append([table, target.id, before, {k: _jsonable(v) for k, v in changes.items()}])


async def _category_paths(db, org_id, entity_type) -> dict:
    cats = (await db.execute(select(Category).where(
        Category.org_id == org_id, Category.entity_type == entity_type))).scalars().all()
    by_id = {c.id: c for c in cats}

    def path(c):
        names, seen = [], set()
        while c and c.id not in seen:
            seen.add(c.id)
            names.append(c.name)
            c = by_id.get(c.parent_id)
        return " > ".join(reversed(names))
    return {_path_key(path(c)): (c.id, path(c), c.name) for c in cats}


def _path_key(p: str) -> str:
    return " > ".join(x.strip().lower() for x in str(p).split(">") if x.strip())


async def _ensure_category(ctx: Ctx, entity_type: str, path: str) -> str:
    paths = ctx.data["cats"]
    parent, built = None, []
    for name in [x.strip() for x in path.split(">") if x.strip()]:
        built.append(name)
        k = _path_key(" > ".join(built))
        if k not in paths:
            c = Category(org_id=ctx.org.id, entity_type=entity_type, name=name[:120], parent_id=parent)
            ctx.db.add(c)
            await ctx.db.flush()
            paths[k] = (c.id, " > ".join(built), c.name)
            ctx.undo["created"].append(["categories", c.id])
        parent = paths[k][0]
    return parent


def _plan_category(ctx: Ctx, row: Row, field_key="category"):
    path = row.rec.get(field_key)
    if not path:
        return
    hit = ctx.data["cats"].get(_path_key(path))
    if hit:
        row.extra["category_id"], row.extra["category_name"] = hit[0], hit[2]
    elif ctx.options.get("create_categories"):
        row.extra["new_category"] = path
        row.warnings.append(f"Category '{path}' will be created")
    else:
        row.errors.append(f"Category '{path}' doesn't exist — add it first or choose “Create missing categories”")


# ── kinds ────────────────────────────────────────────────────────────────────
class Kind:
    key = ""
    sheet = ""
    noun = "record"
    table = ""
    access: tuple = ()          # any one of these opens the importer
    fields: list = []
    match_help = ""

    def field(self, key) -> F:
        return next(f for f in self.fields if f.key == key)

    async def load(self, ctx: Ctx):
        pass

    async def plan(self, ctx: Ctx, row: Row):
        raise NotImplementedError

    async def apply(self, ctx: Ctx, row: Row):
        raise NotImplementedError

    async def lists(self, db, org_id) -> dict:
        return {}


UNITS = ("pcs", "m", "kg", "l", "hr", "set")


class Products(Kind):
    key, sheet, noun, table = "products", "Products", "product", "parts"
    access = ("can_create_products", "can_edit_products", "can_manage_inventory")
    fields = [
        F("name", "Name", required=True, maxlen=300, synonyms=("product", "product name", "item name", "title", "description name"),
          help="Product name as it appears on jobs and receipts", example="iPhone 13 Screen (OLED)"),
        F("sku", "SKU", "code", maxlen=100, synonyms=("item", "item #", "item no", "part", "part no", "part number", "article", "art no"),
          help="Your stock-keeping code. Kept as text, so leading zeros stay", example="SCR-IP13-OLED"),
        F("barcode", "Barcode", "barcode", maxlen=100, synonyms=("ean", "upc", "gtin", "ean13", "barcode number"),
          help="EAN / UPC / any code your scanner reads. Kept as text", example="0012345678905"),
        F("code", "Code", "code", maxlen=100, synonyms=("internal code", "product code"), help="Internal article code (optional)"),
        F("mpn", "MPN", "code", maxlen=64, synonyms=("manufacturer part number", "mfr part", "mfr pn", "oem part"),
          help="Manufacturer part number — one product per MPN"),
        F("category", "Category", list_name="categories", maxlen=250, synonyms=("group", "product category", "category path"),
          help="Pick from the list. Sub-categories as Parent > Child", example=""),
        F("product_type", "Product type", "choice", choices=("Standard", "Serialized"), synonyms=("type", "tracking"),
          help="Serialized = every unit has its own serial / IMEI. Can't be turned off later", example="Standard"),
        F("serialized", "Serialized", "bool", in_template=False, synonyms=("serial tracked", "track serials")),
        F("expiry_tracking", "Track expiry", "bool", synonyms=("expiry", "expiry tracking"),
          help="yes = batches carry an expiry date (used first-expiring-first)", example="no"),
        F("uom", "Unit", "choice", choices=UNITS, synonyms=("uom", "unit of measure", "units"), help="pcs, m, kg, l, hr or set", example="pcs"),
        F("cost_price", "Cost price", "num", cost=True, synonyms=("cost", "purchase price", "buy price", "unit cost"),
          help="What you pay the supplier", example=40),
        F("sell_price", "Sell price", "num", synonyms=("price", "retail price", "sale price", "retail"), help="Counter / POS price", example=99),
        F("repair_price", "Repair price", "num", synonyms=("job price", "price on jobs"),
          help="Price when used on a job (blank = sell price)"),
        F("min_stock", "Min stock", "int", synonyms=("minimum stock", "reorder level", "reorder point", "min qty"),
          help="Low-stock alert below this", example=2),
        F("warranty_days", "Warranty (days)", "int", synonyms=("client warranty days", "warranty", "client warranty"), example=90),
        F("supplier_warranty_days", "Supplier warranty (days)", "int", synonyms=("supplier warranty",)),
        F("brand", "Brand", maxlen=100, synonyms=("manufacturer", "make")),
        F("description", "Description", maxlen=5000, synonyms=("notes", "details")),
        F("opening_qty", "Opening stock", "num", synonyms=("qty", "quantity", "stock", "on hand", "in stock", "stock qty"),
          help="Units to post into stock now (added to what's there). Leave blank for none", example=5),
        F("bin", "Bin", maxlen=60, synonyms=("bin location", "shelf", "location bin"), help="Bin name in the chosen warehouse (optional)"),
        F("serials", "Serial numbers", "serials", synonyms=("serials", "serial", "imei", "imeis", "serial no", "sn"),
          help="For serialized products: one per unit, separated by commas. Missing ones can be generated"),
    ]
    match_help = "SKU, then barcode, then name (or the identifier chosen in Settings → Inventory)"

    async def lists(self, db, org_id):
        paths = await _category_paths(db, org_id, "product")
        return {"categories": sorted(v[1] for v in paths.values()), "product types": list(self.field("product_type").choices)}

    async def load(self, ctx):
        db, org = ctx.db, ctx.org
        key = (org.inv_unique_identifier or "").lower()
        ctx.data["keys"] = [key] if key in ("sku", "barcode", "code", "name") else ["sku", "barcode", "name"]
        parts = (await db.execute(select(Part).where(Part.org_id == org.id, Part.active == True))).scalars().all()  # noqa: E712
        maps = {k: {} for k in ("sku", "barcode", "code", "name", "mpn")}
        for p in parts:
            for k in ("sku", "barcode", "code", "mpn"):
                v = getattr(p, k)
                if v:
                    maps[k].setdefault(self._norm(k, v), []).append(p)
            maps["name"].setdefault(p.name.strip().lower(), []).append(p)
        ctx.data["maps"] = maps
        ctx.data["cats"] = await _category_paths(db, org.id, "product")
        loc_id = ctx.options.get("location_id")
        q = select(Location).where(Location.org_id == org.id, Location.active == True)  # noqa: E712
        loc = (await db.execute(q.where(Location.id == loc_id))).scalar_one_or_none() if loc_id else \
            (await db.execute(q.order_by(Location.created_at, Location.id).limit(1))).scalar_one_or_none()
        if loc_id and not loc:
            raise HTTPException(404, "Warehouse not found")
        ctx.data["location"] = loc
        ctx.data["location_denied"] = None
        if loc is not None:
            from backend.services import warehouses as wh
            try:
                await wh.guard_action(db, ctx.current, "post", location_id=loc.id)
            except HTTPException as e:
                ctx.data["location_denied"] = e.detail
        bins = (await db.execute(select(StorageBin).where(
            StorageBin.org_id == org.id, StorageBin.active == True,  # noqa: E712
            StorageBin.location_id == (loc.id if loc else None)))).scalars().all()
        ctx.data["bins"] = {b.name.strip().lower(): b for b in bins}
        ctx.data["serials_in_stock"] = set((await db.execute(select(SerializedItem.serial_number).where(
            SerializedItem.org_id == org.id,
            SerializedItem.status.in_(("in_stock", "allocated_to_order"))))).scalars())
        ctx.data["serial_n"] = (await db.execute(select(func.count()).select_from(SerializedItem).where(
            SerializedItem.org_id == org.id))).scalar() or 0

    @staticmethod
    def _norm(k, v):
        return C.code_key(v) if k == "barcode" else str(v).strip().lower()

    async def plan(self, ctx, row):
        rec, maps = row.rec, ctx.data["maps"]
        label = {"sku": "SKU", "barcode": "Barcode", "code": "Code", "name": "Name"}
        for k in ctx.data["keys"]:
            v = rec.get(k)
            if not v:
                continue
            hits = maps[k].get(self._norm(k, v), [])
            if len(hits) > 1:
                row.errors.append(f"{label[k]} {v} matches {len(hits)} products — make it unique first")
                return
            if hits:
                row.target, row.matched_on = hits[0], k
            break
        ident = next(((k, rec[k]) for k in ctx.data["keys"] if rec.get(k)), None)
        if row.target is not None:
            ctx.claim("part", row.target.id, row, f"{label[row.matched_on]} {rec[row.matched_on]}")
        elif ident and ident[0] == "name":
            ctx.claim("u-name", self._norm(*ident), row, f"Name {ident[1]}")
        # codes that identify a product must not land on a second one, in the shop or in this file.
        # The SKU always does (ii144a: one active product per SKU, enforced by the database), whatever
        # the shop's matching key; so does the MPN.
        ident_keys = [k for k in ctx.data["keys"] if k != "name"]
        for k in ident_keys + [k for k in ("sku", "mpn") if k not in ident_keys]:
            v = rec.get(k)
            if not v:
                continue
            shown = f"{label.get(k, 'MPN')} {v}"
            if k != row.matched_on and not ctx.claim(f"u-{k}", self._norm(k, v), row, shown):
                continue
            other = [p for p in maps[k].get(self._norm(k, v), []) if p is not row.target]
            if other:
                row.errors.append(f"{shown} is already used by {other[0].name}")
        row.label = rec.get("name") or (row.target.name if row.target is not None else "")
        _plan_category(ctx, row)

        values = {k: rec[k] for k in ("name", "sku", "barcode", "code", "mpn", "uom", "cost_price", "sell_price",
                                      "repair_price", "min_stock", "warranty_days", "supplier_warranty_days",
                                      "brand", "description", "expiry_tracking") if k in rec}
        for k in ("cost_price", "sell_price", "repair_price"):
            if k in values:
                values[k] = q2(values[k])
        if "category_id" in row.extra:
            values["category_id"] = row.extra["category_id"]
        serialized = rec.get("serialized")
        if serialized is None and "product_type" in rec:
            serialized = rec["product_type"] == "Serialized"
        serials = list(rec.get("serials") or [])
        if serials and serialized is None and row.target is None:
            serialized = True
        if serialized is not None:
            values["serialized"] = serialized

        if row.target is None:
            row.action = "create"
            if not rec.get("name"):
                row.errors.append("Name is required for a new product")
            if not ctx.perms["can_create_products"]:
                row.errors.append("Your role can't create products")
            row.changes = values
            is_serial = bool(serialized)
        else:
            row.action = "update"
            t = row.target
            if "barcode" in values and C.code_key(values["barcode"]) == C.code_key(t.barcode):
                values.pop("barcode")       # same code, Excel only dropped its zeros
            if t.serialized and values.get("serialized") is False:
                values.pop("serialized")
                row.warnings.append("Serial tracking can't be turned off — left on")
            if values.get("serialized") and not t.serialized and not ctx.perms["can_track_serials"]:
                row.errors.append("Your role can't turn on serial tracking")
            row.changes = _changes(t, values)
            if row.changes and not ctx.perms["can_edit_products"]:
                row.errors.append("Your role can't edit products")
            is_serial = t.serialized or bool(values.get("serialized"))
        for k in ("barcode",):          # a shared barcode is allowed, but worth a look
            if rec.get(k) and k not in ctx.data["keys"]:
                other = [p for p in maps[k].get(self._norm(k, rec[k]), []) if p is not row.target]
                if other:
                    row.warnings.append(f"{k.upper() if k == 'sku' else 'Barcode'} {rec[k]} is also on {other[0].name}")
        self._plan_stock(ctx, row, is_serial, serials)
        if row.target is not None and row.extra.get("stock"):
            row.warnings.append(f"Adds {row.extra['stock']['qty']:g} to the stock already there")
        if row.action == "update" and not row.changes and not row.extra.get("stock") \
                and not row.extra.get("new_category") and not row.errors:
            row.action = "skip"
            row.warnings.append("No changes")

    def _plan_stock(self, ctx, row, is_serial, serials):
        rec = row.rec
        qty = rec.get("opening_qty")
        if serials and qty is None:
            qty = float(len(serials))
        if rec.get("bin") and not qty:
            row.warnings.append("Bin ignored — no opening stock on this row")
        if not qty:
            if serials:
                row.errors.append("Serial numbers need an opening stock quantity")
            return
        if qty > MAX_OPENING_QTY:
            row.errors.append("Opening stock looks like a barcode or serial, not a count")
            return
        if not ctx.perms["can_manage_inventory"]:
            row.errors.append("Your role can't post stock (Opening stock) — leave the column blank or ask a manager")
            return
        loc = ctx.data["location"]
        if loc is None:
            row.errors.append("Add a warehouse before posting opening stock")
            return
        if ctx.data["location_denied"]:
            row.errors.append(f"Opening stock: {ctx.data['location_denied']}")
            return
        bin_obj = None
        if rec.get("bin"):
            bin_obj = ctx.data["bins"].get(rec["bin"].strip().lower())
            if bin_obj is None:
                row.errors.append(f"Bin '{rec['bin']}' not found in {loc.name} — add it in Settings → Bins first")
        if is_serial:
            if not float(qty).is_integer():
                row.errors.append("Serialized stock must be a whole number of units")
                return
            n = int(qty)
            if n > MAX_SERIALS:
                row.errors.append(f"At most {MAX_SERIALS:,} serialized units per row")
                return
            if len(serials) > n:
                row.errors.append(f"More serial numbers ({len(serials)}) than units ({n})")
            seen_here = set()
            for sn in serials:
                if sn in seen_here:
                    row.errors.append(f"Serial {sn} is listed twice")
                    continue
                seen_here.add(sn)
                if sn in ctx.data["serials_in_stock"]:
                    row.errors.append(f"Serial {sn} is already in stock")
                else:
                    ctx.claim("serial", sn, row, f"Serial {sn}")
            missing = n - len(serials)
            if missing > 0:
                if ctx.options.get("generate_serials"):
                    try:
                        serials = serials + generate_serials(ctx, missing, row)
                    except HTTPException as e:
                        row.errors.append(e.detail)
                        return
                    row.warnings.append(f"{missing} serial number{'s' if missing != 1 else ''} will be generated")
                else:
                    row.errors.append(f"{n} units need {n} serial numbers ({len(serials)} given) — add them or "
                                      f"choose “Generate missing serials”")
        elif serials:
            row.errors.append("Serial numbers need a serialized product — set Product type to Serialized")
        row.extra["stock"] = {"qty": float(qty), "bin_id": bin_obj.id if bin_obj else None, "serials": serials}

    async def apply(self, ctx, row):
        db = ctx.db
        values = dict(row.changes)
        if row.extra.get("new_category"):
            values["category_id"] = await _ensure_category(ctx, "product", row.extra["new_category"])
        if "category_id" in values:   # the name the POS groups by moves with it (and is undone with it)
            values["category"] = row.extra.get("category_name") or row.extra["new_category"].split(">")[-1].strip()[:100]
        stock = row.extra.get("stock")
        if row.target is None:
            part = Part(org_id=ctx.org.id, **values)
            db.add(part)
            await db.flush()
            ctx.undo["created"].append(["parts", part.id])
            extra_before = None
        else:
            part = row.target
            extra_before = {"avg_cost": _jsonable(part.avg_cost), "cost_price": _jsonable(part.cost_price)} if stock else None
            _apply_update(ctx, "parts", part, values, extra_before)
        if "category_id" in values:
            from backend.routers.catalog_products import _sync_part_category_name
            await _sync_part_category_name(db, ctx.org.id, part)
        if stock:
            cost = row.rec.get("cost_price")
            if cost is None:      # no cost column (or the role can't see costs): keep the average where it is
                cost = part.avg_cost or part.cost_price or 0
            await post_stock(ctx, part, stock, unit_cost=float(cost))
        await db.flush()


async def post_stock(ctx: Ctx, part: Part, stock: dict, unit_cost: float):
    """Receive opening stock through the one stock-inflow engine, tagged with the import batch."""
    from backend.services import stock_posting
    db, loc = ctx.db, ctx.data["location"]
    before = set((await db.execute(select(StockItemBatch.id).where(
        StockItemBatch.part_id == part.id, StockItemBatch.location_id == loc.id))).scalars())
    await stock_posting.apply_receipt_line(
        db, org_id=ctx.org.id, part=part, location_id=loc.id, qty=stock["qty"], unit_cost=unit_cost,
        created_by_id=ctx.current.id, ref_type="import", ref_id=ctx.batch_id, bin_id=stock.get("bin_id"),
        serials=stock.get("serials") or None, create_batch=True)
    await db.flush()
    after = set((await db.execute(select(StockItemBatch.id).where(
        StockItemBatch.part_id == part.id, StockItemBatch.location_id == loc.id))).scalars())
    ctx.undo["stock"].append({"part_id": part.id, "location_id": loc.id, "qty": stock["qty"],
                              "batch_ids": sorted(after - before)})


def serial_template(org) -> str:
    """The shop's product serial template if it has a usable counter ({C} or {C:n}, n ≤ 12)."""
    import re
    tpl = (org.inv_product_serial_template or "").strip()
    m = re.search(r"\{C(?::(\d+))?\}", tpl)
    if not m or (m.group(1) and int(m.group(1)) > 12):
        return "SN-{C:6}"
    return tpl


def generate_serials(ctx: Ctx, n: int, row: Row) -> list[str]:
    """`n` new serials from the shop's product serial template (Settings → Inventory, e.g. SN-{C:6}),
    skipping any already used in stock or elsewhere in this file."""
    from backend.services.numbering import render
    used = ctx.data.get("serials_made", 0) + n
    if used > MAX_SERIALS:
        raise HTTPException(422, f"At most {MAX_SERIALS:,} serial numbers can be generated at a time")
    ctx.data["serials_made"] = used
    tpl = serial_template(ctx.org)
    taken = ctx.data["serials_in_stock"]
    out, tries = [], 0
    while len(out) < n:
        tries += 1
        if tries > n + 100_000:
            raise HTTPException(422, "Couldn't generate unique serial numbers — check the serial template in Settings")
        ctx.data["serial_n"] += 1
        sn = render(tpl, ctx.data["serial_n"])
        if sn in taken or ("serial", sn) in ctx.seen:
            continue
        ctx.seen[("serial", sn)] = row.row
        out.append(sn)
    return out


class Services(Kind):
    key, sheet, noun, table = "services", "Services", "service", "services"
    access = ("can_edit_labors",)
    fields = [
        F("name", "Name", required=True, maxlen=300, synonyms=("service", "service name", "labor", "title"),
          help="Service or labour name", example="Screen replacement — iPhone 13"),
        F("is_labor", "Type", "choice", choices=("Service", "Labor"), synonyms=("service type", "kind"),
          help="Service or Labor", example="Service"),
        F("category", "Category", list_name="categories", maxlen=250, synonyms=("group",), help="Pick from the list"),
        F("unit", "Unit", maxlen=24, synonyms=("uom",), help="pc, hr, …", example="pc"),
        F("price", "Price", "num", synonyms=("sell price", "rate", "retail price"), example=89),
        F("cost_price", "Cost price", "num", cost=True, synonyms=("cost",), example=0),
        F("duration_min", "Duration (min)", "int", synonyms=("duration", "minutes", "time"), example=60),
        F("warranty_days", "Warranty (days)", "int", synonyms=("warranty",), example=90),
        F("barcode", "Barcode", "code", maxlen=100, synonyms=("code", "sku")),
        F("bookable", "Bookable online", "bool", synonyms=("bookable", "online booking"), example="yes"),
        F("flat_rate", "Flat-rate labour", "bool", synonyms=("flat rate",), help="yes = billed as booked hours × price"),
        F("description", "Description", maxlen=2000, synonyms=("notes", "details")),
    ]
    match_help = "barcode, then name"

    async def lists(self, db, org_id):
        paths = await _category_paths(db, org_id, "service")
        return {"categories": sorted(v[1] for v in paths.values())}

    async def load(self, ctx):
        rows = (await ctx.db.execute(select(Service).where(
            Service.org_id == ctx.org.id, Service.active == True))).scalars().all()  # noqa: E712
        ctx.data["by_barcode"] = {C.code_key(s.barcode): s for s in rows if s.barcode}
        ctx.data["by_name"] = {s.name.strip().lower(): s for s in rows}
        ctx.data["cats"] = await _category_paths(ctx.db, ctx.org.id, "service")

    async def plan(self, ctx, row):
        rec = row.rec
        if rec.get("barcode") and C.code_key(rec["barcode"]) in ctx.data["by_barcode"]:
            row.target, row.matched_on = ctx.data["by_barcode"][C.code_key(rec["barcode"])], "barcode"
        elif rec.get("name") and rec["name"].lower() in ctx.data["by_name"]:
            row.target, row.matched_on = ctx.data["by_name"][rec["name"].lower()], "name"
        row.label = rec.get("name") or (row.target.name if row.target is not None else "")
        if row.target is not None:
            ctx.claim("service", row.target.id, row, row.label)
        elif rec.get("name"):
            ctx.claim("service-name", rec["name"].lower(), row, f"Name {rec['name']}")
        _plan_category(ctx, row)
        values = {k: rec[k] for k in ("name", "unit", "price", "cost_price", "duration_min", "warranty_days",
                                      "barcode", "bookable", "flat_rate", "description") if k in rec}
        if "is_labor" in rec:
            values["is_labor"] = rec["is_labor"] == "Labor"
        if "category_id" in row.extra:
            values["category_id"], values["category"] = row.extra["category_id"], row.extra["category_name"]
        if row.target is None:
            row.action = "create"
            if not rec.get("name"):
                row.errors.append("Name is required for a new service")
            row.changes = values
        else:
            row.action = "update"
            row.changes = _changes(row.target, values)

    async def apply(self, ctx, row):
        values = dict(row.changes)
        if row.extra.get("new_category"):
            values["category_id"] = await _ensure_category(ctx, "service", row.extra["new_category"])
            values["category"] = row.extra["new_category"].split(">")[-1].strip()
        if row.target is None:
            s = Service(org_id=ctx.org.id, **values)
            ctx.db.add(s)
            await ctx.db.flush()
            ctx.undo["created"].append(["services", s.id])
        else:
            _apply_update(ctx, "services", row.target, values)


class Clients(Kind):
    key, sheet, noun, table = "clients", "Clients", "client", "clients"
    access = ("can_create_clients",)
    fields = [
        F("name", "Name", required=True, maxlen=200, synonyms=("full name", "client", "customer", "contact", "company", "client name"),
          help="Person or company name", example="Jane Smith"),
        F("client_type", "Type", "choice", choices=("Person", "Organization"), synonyms=("client type", "contact type"),
          help="Person or Organization", example="Person"),
        F("phone", "Phone", "code", maxlen=50, synonyms=("mobile", "telephone", "phone number", "cell", "tel"),
          help="Kept as text so the leading 0 or + stays. Used to find existing clients", example="+1 555 010 0001"),
        F("email", "Email", maxlen=200, synonyms=("e mail", "email address", "mail"), example="jane@example.com"),
        F("address", "Address", maxlen=1000, synonyms=("street", "postal address")),
        F("notes", "Notes", maxlen=5000, synonyms=("comment", "comments", "note")),
        F("source", "Source", maxlen=100, synonyms=("lead source", "how did you hear", "referral")),
        F("discount_card", "Discount card", "code", maxlen=60, synonyms=("loyalty card", "card number")),
        F("tax_id", "Tax ID", "code", maxlen=80, synonyms=("vat", "vat number", "tax number", "gst", "ein")),
        F("birthday", "Birthday", "date", synonyms=("date of birth", "dob", "birth date"), help="YYYY-MM-DD"),
        F("consent_email", "Email marketing", "bool", synonyms=("email consent", "accepts email")),
        F("consent_phone", "SMS marketing", "bool", synonyms=("sms consent", "accepts sms", "phone consent")),
    ]
    match_help = "phone, then email"

    async def load(self, ctx):
        rows = (await ctx.db.execute(select(Client).where(Client.org_id == ctx.org.id))).scalars().all()
        ctx.data["by_phone"], ctx.data["by_email"] = {}, {}
        for c in rows:
            if C.digits(c.phone):
                ctx.data["by_phone"].setdefault(C.digits(c.phone), c)
            if c.email:
                ctx.data["by_email"].setdefault(c.email.strip().lower(), c)

    async def plan(self, ctx, row):
        rec = row.rec
        ph, em = C.digits(rec.get("phone")), (rec.get("email") or "").lower()
        if ph and ph in ctx.data["by_phone"]:
            row.target, row.matched_on = ctx.data["by_phone"][ph], "phone"
        elif em and em in ctx.data["by_email"]:
            row.target, row.matched_on = ctx.data["by_email"][em], "email"
        if row.target is not None:
            ctx.claim("client", row.target.id, row, rec.get(row.matched_on))
        elif ph:
            ctx.claim("phone", ph, row, f"Phone {rec['phone']}")
        elif em:
            ctx.claim("email", em, row, f"Email {rec['email']}")
        row.label = rec.get("name") or (row.target.name if row.target is not None else "")
        values = {k: rec[k] for k in ("name", "phone", "email", "address", "notes", "source", "discount_card",
                                      "tax_id", "consent_email", "consent_phone") if k in rec}
        if "client_type" in rec:
            values["client_type"] = rec["client_type"].lower()
        if "birthday" in rec:
            values["birthday"] = rec["birthday"].isoformat()
        if row.target is None:
            row.action = "create"
            if not rec.get("name"):
                row.errors.append("Name is required for a new client")
            row.changes = values
        else:
            row.action = "update"
            if "phone" in values and C.digits(values["phone"]) == C.digits(row.target.phone):
                values.pop("phone")
            row.changes = _changes(row.target, values)

    async def apply(self, ctx, row):
        if row.target is None:
            c = Client(org_id=ctx.org.id, **row.changes)
            ctx.db.add(c)
            await ctx.db.flush()
            ctx.undo["created"].append(["clients", c.id])
        else:
            _apply_update(ctx, "clients", row.target, row.changes)


class Assets(Kind):
    key, sheet, noun, table = "assets", "Assets", "asset", "devices"
    access = ("can_view_asset_info",)
    fields = [
        F("asset_type", "Type", maxlen=50, synonyms=("asset type", "device type", "kind", "category"),
          help="Device, Vehicle, Tool, … (blank = Device)", example="Device"),
        F("brand", "Brand", required=True, maxlen=100, synonyms=("make", "manufacturer"), example="Apple"),
        F("model", "Model", maxlen=100, synonyms=("device", "model name"), example="iPhone 13"),
        F("serial", "Serial", "code", maxlen=100, synonyms=("serial number", "serial no", "sn"),
          help="Kept as text. Used to find existing assets", example="F2LXK0ABCD12"),
        F("imei", "IMEI", "code", maxlen=40, synonyms=("imei number", "imei1"), help="Kept as text. Used to find existing assets"),
        F("vin", "VIN", "code", maxlen=20, synonyms=("vehicle identification number", "chassis")),
        F("color", "Color", maxlen=50, synonyms=("colour",)),
        F("condition", "Condition", maxlen=50, synonyms=("state", "grade")),
        F("owner_phone", "Owner phone", "code", maxlen=50, synonyms=("client phone", "customer phone", "phone"),
          help="Phone of an existing client. Blank owner = company asset"),
        F("owner_email", "Owner email", maxlen=200, synonyms=("client email", "customer email", "email"),
          help="Email of an existing client (if no phone)"),
        F("purchase_date", "Purchase date", "date", synonyms=("bought", "purchased", "date purchased"), help="YYYY-MM-DD"),
        F("purchase_price", "Purchase price", "num", cost=True, synonyms=("cost", "price paid")),
    ]
    match_help = "IMEI, then serial"
    cost_perm = "can_see_asset_cost"

    async def load(self, ctx):
        from backend.services import identifiers as ids
        rows = (await ctx.db.execute(select(Device).where(Device.org_id == ctx.org.id))).scalars().all()
        ctx.data["by_imei"] = {d.imei: d for d in rows if d.imei}
        ctx.data["by_serial"] = {d.serial: d for d in rows if d.serial}
        clients = (await ctx.db.execute(select(Client).where(Client.org_id == ctx.org.id))).scalars().all()
        ctx.data["c_phone"] = {C.digits(c.phone): c for c in clients if C.digits(c.phone)}
        ctx.data["c_email"] = {c.email.strip().lower(): c for c in clients if c.email}
        ctx.data["ids"] = ids
        from backend.services import industry as ind
        ctx.data["template_key"] = ind.normalize_template_key(ctx.org.industry_template)

    async def plan(self, ctx, row):
        ids, rec = ctx.data["ids"], row.rec
        imei, serial = ids.normalize("imei", rec.get("imei")), ids.normalize("serial", rec.get("serial"))
        if imei and imei in ctx.data["by_imei"]:
            row.target, row.matched_on = ctx.data["by_imei"][imei], "imei"
        elif serial and serial in ctx.data["by_serial"]:
            row.target, row.matched_on = ctx.data["by_serial"][serial], "serial"
        if row.target is not None:
            ctx.claim("asset", row.target.id, row, rec.get(row.matched_on))
        elif imei:
            ctx.claim("imei", imei, row, f"IMEI {rec['imei']}")
        elif serial:
            ctx.claim("serial", serial, row, f"Serial {rec['serial']}")
        row.label = " ".join(x for x in (rec.get("brand"), rec.get("model")) if x) or rec.get("serial") or rec.get("imei") or \
            (" ".join(x for x in ((row.target.brand or ""), (row.target.model or "")) if x) if row.target is not None else "")
        values = {k: rec[k] for k in ("asset_type", "brand", "model", "color", "condition", "purchase_date",
                                      "purchase_price") if k in rec}
        for k in ("serial", "imei", "vin"):
            if k in rec:
                values[k] = ids.normalize(k, rec[k])
                warn = ids.problem(k, values[k])
                if warn:
                    row.warnings.append(warn)
        owner = None
        if rec.get("owner_phone"):
            owner = ctx.data["c_phone"].get(C.digits(rec["owner_phone"]))
            if owner is None:
                row.errors.append(f"No client with phone {rec['owner_phone']} — import clients first")
        elif rec.get("owner_email"):
            owner = ctx.data["c_email"].get(rec["owner_email"].lower())
            if owner is None:
                row.errors.append(f"No client with email {rec['owner_email']} — import clients first")
        if owner is not None:
            values["client_id"], values["owner_type"] = owner.id, "client"
        if row.target is None:
            row.action = "create"
            if not rec.get("brand"):
                row.errors.append("Brand is required for a new asset")
            values.setdefault("asset_type", "Device")
            values.setdefault("owner_type", "company")
            values["asset_status"] = "in_use" if values["owner_type"] == "client" else "in_stock"
            row.changes = values
        else:
            row.action = "update"
            row.changes = _changes(row.target, values)
            if row.changes and not ctx.perms["can_manage_inventory"]:
                row.errors.append("Your role can't edit existing assets")

    async def apply(self, ctx, row):
        if row.target is None:
            d = Device(org_id=ctx.org.id, template_key=ctx.data["template_key"], **row.changes)
            ctx.db.add(d)
            await ctx.db.flush()
            ctx.undo["created"].append(["devices", d.id])
        else:
            _apply_update(ctx, "devices", row.target, row.changes)


class Suppliers(Kind):
    key, sheet, noun, table = "suppliers", "Suppliers", "supplier", "suppliers"
    access = ("can_manage_inventory",)
    fields = [
        F("name", "Name", required=True, maxlen=200, synonyms=("supplier", "supplier name", "vendor", "company"),
          help="Supplier name. Used to find existing suppliers", example="Acme Parts Ltd"),
        F("contact_name", "Contact person", maxlen=200, synonyms=("contact", "contact name", "rep"), example="Sam Lee"),
        F("phone", "Phone", "code", maxlen=50, synonyms=("telephone", "phone number", "mobile"), example="+1 555 010 0099"),
        F("email", "Email", maxlen=200, synonyms=("e mail", "email address"), example="orders@acme.example"),
        F("lead_days", "Lead time (days)", "int", synonyms=("lead time", "lead days", "delivery days"), example=3),
        F("currency", "Currency", maxlen=8, synonyms=("ccy",)),
    ]
    match_help = "name"

    async def load(self, ctx):
        rows = (await ctx.db.execute(select(Supplier).where(Supplier.org_id == ctx.org.id))).scalars().all()
        ctx.data["by_name"] = {s.name.strip().lower(): s for s in rows}

    async def plan(self, ctx, row):
        rec = row.rec
        name = (rec.get("name") or "").lower()
        if name and name in ctx.data["by_name"]:
            row.target, row.matched_on = ctx.data["by_name"][name], "name"
        if name:
            ctx.claim("supplier", name, row, f"Name {rec['name']}")
        row.label = rec.get("name") or ""
        values = {k: rec[k] for k in ("name", "contact_name", "phone", "email", "lead_days", "currency") if k in rec}
        if row.target is None:
            row.action = "create"
            if not name:
                row.errors.append("Name is required for a new supplier")
            row.changes = values
        else:
            row.action = "update"
            row.changes = _changes(row.target, values)

    async def apply(self, ctx, row):
        db = ctx.db
        if row.target is not None:
            _apply_update(ctx, "suppliers", row.target, row.changes)
            return
        v = row.changes
        terms = [Client.phone == v["phone"]] if v.get("phone") else []
        if v.get("email"):
            terms.append(Client.email == v["email"])
        contact = (await db.execute(select(Client).where(Client.org_id == ctx.org.id, or_(*terms)))).scalars().first() \
            if terms else None
        if contact is None:
            contact = Client(org_id=ctx.org.id, name=v["name"], phone=v.get("phone"), email=v.get("email"), is_supplier=True)
            db.add(contact)
            await db.flush()
            ctx.undo["created"].append(["clients", contact.id])
        elif not contact.is_supplier:
            _apply_update(ctx, "clients", contact, {"is_supplier": True})
        s = Supplier(org_id=ctx.org.id, balance=0, client_id=contact.id, **v)
        db.add(s)
        await db.flush()
        ctx.undo["created"].append(["suppliers", s.id])


KINDS: dict[str, Kind] = {k.key: k for k in (Products(), Services(), Clients(), Assets(), Suppliers())}

# Postings / purchase orders keep their own paste-and-match flow; only their columns are mapped here.
POSTING_FIELDS = [
    F("code", "Code / SKU", synonyms=("code", "sku", "part", "part#", "part no", "partno", "part number", "item", "item #",
                                      "item no", "barcode", "upc", "ref")),
    F("mpn", "MPN", synonyms=("mpn", "mfr pn", "mfr part", "mfg pn", "manufacturer part", "manufacturer part number", "oem part")),
    F("name", "Name", synonyms=("name", "description", "desc", "product", "product name", "item name", "title", "details")),
    F("qty", "Qty", synonyms=("qty", "quantity", "count", "units", "on hand", "received", "qty received", "pcs", "stock")),
    F("cost", "Cost", synonyms=("cost", "unit cost", "price", "unit price", "buy price", "purchase price", "wholesale", "net", "rate")),
    F("markup", "Markup %", synonyms=("markup", "margin", "markup%", "margin%", "uplift")),
    F("serials", "Serials", synonyms=("serials", "serial", "serial numbers", "serial no", "imei", "imeis", "sn")),
]
MAPPING_KINDS = set(KINDS) | {"posting", "po"}


def fields_for(kind: str) -> list[F]:
    return POSTING_FIELDS if kind in ("posting", "po") else KINDS[kind].fields


def access_for(kind: str) -> tuple:
    return ("can_manage_inventory",) if kind in ("posting", "po") else KINDS[kind].access


def suggest_mapping(kind: str, headers: list[str], remembered: Optional[dict]) -> dict:
    fields = fields_for(kind)
    keys = {f.key for f in fields}
    out, used = {}, set()
    skipped = {h for h, f in (remembered or {}).items() if f == "ignore"}   # columns the user chose not to import
    for h, f in (remembered or {}).items():
        if h in headers and f in keys and f not in used:
            out[h], used = f, used | {f}
    for pass_ in ("label", "synonym"):
        for h in headers:
            if h in out or h in skipped:
                continue
            n = C.norm_header(h)
            for f in fields:
                if f.key in used:
                    continue
                names = [C.norm_header(f.label), C.norm_header(f.key.replace("_", " "))] if pass_ == "label" \
                    else [C.norm_header(s) for s in f.synonyms]
                if n in names:
                    out[h] = f.key
                    used.add(f.key)
                    break
    return out


# ── the runner ───────────────────────────────────────────────────────────────
async def perms_for(db, current) -> dict:
    keys = {"can_create_products", "can_edit_products", "can_manage_inventory", "can_track_serials",
            COST_PERM, "can_see_asset_cost"}
    return {k: await has_perm(db, current, k) for k in keys}


async def assert_access(db, current, kind: str):
    if not any([await has_perm(db, current, p) for p in access_for(kind)]):
        raise HTTPException(403, "Your role can't import into this list")


def _is_example(kind: Kind, rec: dict) -> bool:
    ex = {f.key: f.example for f in kind.fields if f.example not in ("", None)}
    if not ex or set(rec) - set(ex) or not rec:
        return False
    types = {f.key: f.type for f in kind.fields}
    for k, v in rec.items():
        e = ex[k]
        if types[k] in ("code", "barcode"):
            if C.code_key(str(v)) != C.code_key(str(e)):     # Excel may have stripped the zeros
                return False
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            if not isinstance(e, (int, float)) or float(v) != float(e):
                return False
        elif isinstance(v, bool):
            if v != (str(e).lower() == "yes"):
                return False
        elif str(v).strip().lower() != str(e).strip().lower():
            return False
    return True


async def run(db, current, kind_key: str, headers: list, rows: list, mapping: dict, options: dict, *,
              commit: bool, file_name: Optional[str] = None, row_numbers: Optional[list] = None) -> dict:
    kind = KINDS[kind_key]
    await assert_access(db, current, kind_key)
    by_key = {f.key: f for f in kind.fields}
    headers = [str(h) for h in headers]
    if not mapping:
        raise HTTPException(422, "Match at least one column to a field")
    if set(mapping.values()) - set(by_key):
        raise HTTPException(422, f"Unknown fields: {sorted(set(mapping.values()) - set(by_key))}")
    if len(set(mapping.values())) != len(mapping):
        raise HTTPException(422, "Each field can only come from one column")
    if set(mapping) - set(headers):
        raise HTTPException(422, "The mapping names a column that isn't in the file")
    perms = await perms_for(db, current)
    perms.update({p: await has_perm(db, current, p) for p in kind.access})
    warnings = []
    cost_perm = getattr(kind, "cost_perm", COST_PERM)
    if not await has_perm(db, current, cost_perm):
        dropped = [h for h, k in mapping.items() if by_key[k].cost]
        if dropped:
            mapping = {h: k for h, k in mapping.items() if h not in dropped}
            warnings.append(f"Cost columns ignored ({', '.join(dropped)}): your role can't see purchase prices")
    org = (await db.execute(select(Organisation).where(Organisation.id == current.org_id))).scalar_one()
    ctx = Ctx(db=db, current=current, org=org, options=options or {}, perms=perms)
    await kind.load(ctx)
    col = {k: headers.index(h) for h, k in mapping.items()}

    planned: list[Row] = []
    for i, cells in enumerate(rows):
        n = row_numbers[i] if row_numbers and i < len(row_numbers) else i + 2
        rec, errs = {}, []
        for k, ci in col.items():
            raw = cells[ci] if ci < len(cells) else None
            raw = "" if raw is None else C._cell_text(raw)
            raw = str(unguard_cell(raw)).strip()
            if raw == "":
                continue
            try:
                v = parse_cell(by_key[k], raw)
            except C.CellError as e:
                errs.append(str(e))
                continue
            if v not in ("", None, []):
                rec[k] = v
        r = Row(row=n, rec=rec, errors=errs)
        if not rec and not errs:
            r.action = "skip"
            r.warnings.append("Empty row")
        elif _is_example(kind, rec):
            r.action = "skip"
            r.label = str(rec.get("name") or "")
            r.warnings.append("Example row from the template — skipped")
        else:
            await kind.plan(ctx, r)
        if r.errors:
            r.action = "error"
        planned.append(r)

    counts = {a: sum(1 for r in planned if r.action == a) for a in ("create", "update", "error", "skip")}
    result = {"kind": kind_key, "counts": counts, "warnings": warnings, "rows": [r.out() for r in planned],
              "match_on": kind.match_help if kind_key != "products" else _key_help(ctx)}
    if not commit:
        await db.rollback()
        return result

    batch = ImportBatch(org_id=org.id, kind=kind_key, file_name=(file_name or "")[:255] or None, actor_id=current.id)
    db.add(batch)
    await db.flush()
    ctx.batch_id = batch.id
    created = updated = 0
    for r in planned:
        if r.action not in ("create", "update"):
            continue
        mark = {k: len(v) for k, v in ctx.undo.items()}
        cats = dict(ctx.data["cats"]) if "cats" in ctx.data else None
        try:
            async with db.begin_nested():
                await kind.apply(ctx, r)
        except Exception as e:     # one row's surprise never sinks the file
            for k, n in mark.items():
                del ctx.undo[k][n:]
            if cats is not None:
                ctx.data["cats"] = cats     # categories made in the rolled-back savepoint are gone
            r.action = "error"
            r.errors.append(getattr(e, "detail", None) or "Couldn't save this row")
            continue
        if r.action == "create":
            created += 1
        else:
            updated += 1
    failed = sum(1 for r in planned if r.action == "error")
    from backend.services import audit
    if created or updated:
        batch.created_count, batch.updated_count, batch.failed_count = created, updated, failed
        batch.undo = ctx.undo
        await audit.record(db, org_id=org.id, actor_id=current.id, action="import", entity_type=kind_key,
                           entity_id=batch.id, diff={"file_name": file_name, "created": created,
                                                     "updated": updated, "failed": failed})
    else:
        await db.delete(batch)
    await db.commit()
    result.update({"batch_id": batch.id if (created or updated) else None, "created": created, "updated": updated,
                   "failed": failed, "rows": [r.out() for r in planned],
                   "counts": {a: sum(1 for r in planned if r.action == a) for a in ("create", "update", "error", "skip")},
                   "undo_until": (batch.created_at + UNDO_WINDOW).isoformat() if (created or updated) else None})
    return result


def _key_help(ctx: Ctx) -> str:
    keys = ctx.data.get("keys") or []
    names = {"sku": "SKU", "barcode": "barcode", "code": "code", "name": "name"}
    return ", then ".join(names[k] for k in keys)


# ── undo ─────────────────────────────────────────────────────────────────────
def _models_by_table() -> dict:
    from backend.db.database import Base
    return {m.class_.__tablename__: m.class_ for m in Base.registry.mappers if hasattr(m.class_, "__tablename__")}


_STOCK_TABLES = {"stock_levels", "stock_movements", "stock_item_batches", "serialized_items"}


async def _references(db, table: str, ids: list, created: dict) -> int:
    """Rows anywhere in the schema that point at these ids — except rows this import created itself."""
    from backend.db.database import Base
    n = 0
    for t in Base.metadata.sorted_tables:
        if table == "parts" and t.name in _STOCK_TABLES:
            continue
        for c in t.columns:
            if any(fk.column.table.name == table and fk.column.name == "id" for fk in c.foreign_keys):
                q = select(func.count()).select_from(t).where(c.in_(ids))
                own = created.get(t.name)
                if own and "id" in t.c:
                    q = q.where(not_(t.c.id.in_(own)))
                n += (await db.execute(q)).scalar() or 0
    return n


def can_undo(batch: ImportBatch, latest_id: Optional[str]) -> bool:
    return (batch.undone_at is None and batch.id == latest_id and batch.undo is not None
            and datetime.utcnow() - batch.created_at <= UNDO_WINDOW)


async def latest_batch_id(db, org_id: str, kind: str) -> Optional[str]:
    return (await db.execute(select(ImportBatch.id).where(ImportBatch.org_id == org_id, ImportBatch.kind == kind)
                             .order_by(ImportBatch.created_at.desc(), ImportBatch.id.desc()).limit(1))).scalar()


async def undo(db, current, batch: ImportBatch) -> dict:
    if batch.undone_at is not None:
        raise HTTPException(409, "This import was already undone")
    if batch.id != await latest_batch_id(db, batch.org_id, batch.kind):
        raise HTTPException(409, "Only the most recent import of this list can be undone")
    if datetime.utcnow() - batch.created_at > UNDO_WINDOW:
        raise HTTPException(409, "Imports can only be undone within 24 hours")
    u = batch.undo or {}
    models = _models_by_table()
    created: dict = {}
    for t, i in u.get("created", []):
        created.setdefault(t, []).append(i)
    stock = u.get("stock", [])

    used = 0
    for t, ids in created.items():
        if t != "categories":      # a category made by the import is removed only if nothing else uses it
            used += await _references(db, t, ids, created)
    stocked = sorted({s["part_id"] for s in stock} - set(created.get("parts", [])))
    if stocked:       # an existing part got opening stock: any other stock move since makes the old cost stale
        used += (await db.execute(select(func.count()).select_from(StockMovement).where(
            StockMovement.part_id.in_(stocked), StockMovement.created_at >= batch.created_at,
            not_(and_(StockMovement.ref_type == "import", StockMovement.ref_id == batch.id))))).scalar() or 0
    if created.get("parts"):
        used += (await db.execute(select(func.count()).select_from(StockMovement).where(
            StockMovement.part_id.in_(created["parts"]),
            not_(and_(StockMovement.ref_type == "import", StockMovement.ref_id == batch.id))))).scalar() or 0
    all_batches = [b for s in stock for b in s["batch_ids"]]
    if all_batches:
        bs = (await db.execute(select(StockItemBatch).where(StockItemBatch.id.in_(all_batches)))).scalars().all()
        used += sum(1 for b in bs if Decimal(str(b.qty_remaining)) != Decimal(str(b.qty_received)))
        used += (await db.execute(select(func.count()).select_from(SerializedItem).where(
            SerializedItem.batch_id.in_(all_batches), SerializedItem.status != "in_stock"))).scalar() or 0
    if used:
        raise HTTPException(409, f"Can't undo: {used} of the imported records or stock have been used since "
                                 f"(on jobs, sales, documents or stock moves)")
    # claim the undo first, so two clicks (or two tabs) can't both reverse the stock
    claimed = await db.execute(update(ImportBatch).where(ImportBatch.id == batch.id, ImportBatch.undone_at.is_(None))
                               .values(undone_at=datetime.utcnow()))
    if claimed.rowcount != 1:
        await db.rollback()
        raise HTTPException(409, "This import was already undone")

    for s in stock:
        await db.execute(delete(SerializedItem).where(SerializedItem.batch_id.in_(s["batch_ids"])))
        await db.execute(delete(StockItemBatch).where(StockItemBatch.id.in_(s["batch_ids"])))
        await db.execute(delete(StockMovement).where(
            StockMovement.part_id == s["part_id"], StockMovement.ref_type == "import", StockMovement.ref_id == batch.id))
        sl = (await db.execute(select(StockLevel).where(
            StockLevel.part_id == s["part_id"], StockLevel.location_id == s["location_id"]))).scalars().first()
        if sl is not None:
            sl.qty_available = (sl.qty_available or 0) - s["qty"]
    kept = 0
    for entry in reversed(u.get("updated", [])):
        t, i, before = entry[0], entry[1], entry[2]
        after = entry[3] if len(entry) > 3 else {}
        obj = await db.get(models[t], i)
        if obj is None:
            continue
        for k, v in before.items():
            if k in after and not _same(_jsonable(getattr(obj, k)), after[k]):
                kept += 1         # changed again after the import: the newer value wins
                continue
            col = models[t].__table__.c.get(k)
            if isinstance(v, str) and col is not None and isinstance(col.type, Date):
                v = date.fromisoformat(v[:10])
            setattr(obj, k, v)
    await db.flush()
    kept_categories = 0
    for t, i in reversed(u.get("created", [])):
        if t == "categories" and await _references(db, "categories", [i], {}):
            kept_categories += 1
            continue
        if t == "parts":
            for m in (StockLevel, StockMovement, StockItemBatch, SerializedItem):
                await db.execute(delete(m).where(m.part_id == i))
        await db.execute(delete(models[t]).where(models[t].id == i))
    from backend.services import audit
    await audit.record(db, org_id=batch.org_id, actor_id=current.id, action="import_undo", entity_type=batch.kind,
                       entity_id=batch.id, diff={"file_name": batch.file_name, "removed": len(u.get("created", [])),
                                                 "restored": len(u.get("updated", []))})
    await db.commit()
    return {"ok": True, "removed": len(u.get("created", [])), "restored": len(u.get("updated", [])),
            "kept_newer": kept}


# ── templates ────────────────────────────────────────────────────────────────
def _example_text(f: F) -> str:
    return "" if f.example in ("", None) else str(f.example)


async def template(db, org_id: str, kind_key: str, fmt: str) -> tuple[bytes, str, str]:
    kind = KINDS[kind_key]
    cols = [f for f in kind.fields if f.in_template]
    labels = [f.label + (" *" if f.required else "") for f in cols]
    if fmt == "csv":
        import csv as _csv
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerow(labels)
        w.writerow([_example_text(f) for f in cols])
        return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv", f"tekly-{kind_key}-template.csv"

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    lists = await kind.lists(db, org_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = kind.sheet
    ws.append(labels)
    head_fill, req_fill = PatternFill("solid", fgColor="E8EEF9"), PatternFill("solid", fgColor="FDE7D8")
    for i, f in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = req_fill if f.required else head_fill
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(f.label) + 6))
    ws.freeze_panes = "A2"
    lsheet = None
    list_cols = {}
    if lists:
        lsheet = wb.create_sheet("Lists")
        for j, (name, values) in enumerate(lists.items(), start=1):
            lsheet.cell(row=1, column=j, value=name.capitalize()).font = Font(bold=True)
            for r, v in enumerate(values, start=2):
                lsheet.cell(row=r, column=j, value=safe_cell(v))
            list_cols[name] = (get_column_letter(j), max(2, len(values) + 1))
    for i, f in enumerate(cols, start=1):
        letter = get_column_letter(i)
        rng = f"{letter}2:{letter}5000"
        formula = None
        if f.type == "bool":
            formula = '"yes,no"'
        elif f.type == "choice":
            if f.key == "product_type" and "product types" in list_cols:
                lc, last = list_cols["product types"]
                formula = f"Lists!${lc}$2:${lc}${last}"
            else:
                formula = '"' + ",".join(f.choices) + '"'
        elif f.list_name and f.list_name in list_cols:
            lc, last = list_cols[f.list_name]
            formula = f"Lists!${lc}$2:${lc}${last}"
        if formula:
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=f.type != "text")
            dv.error, dv.errorTitle = f"Pick a value from the list for {f.label}", f.label
            ws.add_data_validation(dv)
            dv.add(rng)
        if f.type in ("code", "barcode"):
            ws.column_dimensions[letter].number_format = "@"
        elif f.type == "date":
            ws.column_dimensions[letter].number_format = "yyyy-mm-dd"
    ex = []
    for f in cols:
        v = f.example
        ex.append(str(v) if f.type in ("code", "barcode") and v not in ("", None) else (None if v == "" else v))
    ws.append([safe_cell(v) for v in ex])
    for i, f in enumerate(cols, start=1):
        c = ws.cell(row=2, column=i)
        c.font = Font(italic=True, color="64748B")
        if f.type in ("code", "barcode"):
            c.number_format = "@"

    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 26, 12
    info.column_dimensions["C"].width, info.column_dimensions["D"].width = 70, 28
    info.append([f"Tekly import template — {kind.sheet}"])
    info["A1"].font = Font(bold=True, size=14)
    for line in (
        f"1. Fill one {kind.noun} per row on the “{kind.sheet}” sheet. Columns marked * are required for new {kind.noun}s.",
        "2. The grey row is an example — overwrite it or leave it; it is skipped on import.",
        f"3. Existing {kind.noun}s are found by {kind.match_help} and updated; blank cells never clear a value.",
        "4. Code columns (SKU, barcode, phone) are formatted as text so leading zeros stay.",
        "5. Save as .xlsx (or CSV) and upload it with the Import button. You'll see every row before anything is saved.",
    ):
        info.append([line])
    info.append([])
    info.append(["Column", "Required", "What to enter", "Example"])
    for c in info[info.max_row]:
        c.font = Font(bold=True)
    for f in cols:
        info.append([f.label, "yes" if f.required else "", f.help or "", _example_text(f)])
    for row in info.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return (buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"tekly-{kind_key}-template.xlsx")
