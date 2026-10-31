"""Catalog master (Phase F) — the My Company > Products surface.

Holds the product *definitions* (cards, variants, categories) and catalog-side
operations: serialize toggle (irreversible, Orderry rule), Orderry-style
write-off (serial -> specific serials, non-serial -> qty with FEFO when expiry
tracking is on), barcode generation, and Excel import/export. Stock *balances*
and movements continue to live under /api/inventory.
"""
import asyncio
import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.db.database import get_db
from backend.models import (
    Part, StockLevel, StockMovement, StockItemBatch, SerializedItem,
    Category, Employee,
)
from backend.services.auth import assert_perm, get_current_employee, require_perm
from backend.routers.forms import resolve_custom_data
from backend.services.codes import gtin_check_digit
from backend.services.money import q2
from backend.services.redact import can_see_costs, drop_cost_fields
from backend.core.likes import contains as _contains
from backend.core.spreadsheet import (
    MAX_IMPORT_ROWS, MAX_XLSX_BYTES, open_xlsx, safe_row, too_many_rows, unguard_cell,
)

router = APIRouter(prefix="/api/catalog", tags=["catalog-products"])


# ── Categories (shared tree: product|service|bundle) ─────────────────────────
async def _category_counts(db, org_id, entity_type="product"):
    if entity_type == "bundle":
        from backend.models import Bundle
        rows = (await db.execute(select(Bundle.category_id, func.count()).where(
            Bundle.org_id == org_id, Bundle.category_id.isnot(None), Bundle.active == True)  # noqa: E712
            .group_by(Bundle.category_id))).all()
        return {cid: n for cid, n in rows}
    if entity_type == "service":
        from backend.models import Service
        rows = (await db.execute(select(Service.category_id, func.count()).where(
            Service.org_id == org_id, Service.category_id.isnot(None), Service.active == True)  # noqa: E712
            .group_by(Service.category_id))).all()
        return {cid: n for cid, n in rows}
    rows = (await db.execute(select(Part.category_id, func.count()).where(
        Part.org_id == org_id, Part.category_id.isnot(None), Part.active == True)  # noqa: E712
        .group_by(Part.category_id))).all()
    return {cid: n for cid, n in rows}


async def _category_tree(db, org_id, entity_type, template_key=None):
    rows = (await db.execute(select(Category).where(
        Category.org_id == org_id, Category.entity_type == entity_type))).scalars().all()
    if template_key:
        # a template sees its own categories plus the shared (null-template) ones
        rows = [c for c in rows if (c.template_key or None) in (None, template_key)]
    counts = await _category_counts(db, org_id, entity_type)
    by_parent: dict = {}
    for c in rows:
        by_parent.setdefault(c.parent_id, []).append(c)

    def build(pid):
        out = []
        for c in sorted(by_parent.get(pid, []), key=lambda x: x.name.lower()):
            kids = build(c.id)
            own = counts.get(c.id, 0)
            out.append({"id": c.id, "name": c.name, "auto_deduct": c.auto_deduct,
                        "template_key": c.template_key, "own_count": own,
                        "count": own + sum(k["count"] for k in kids), "children": kids})
        return out
    return build(None)


@router.get("/categories")
async def list_categories(entity_type: str = Query("product"),
                          template_key: Optional[str] = None,
                          db: AsyncSession = Depends(get_db),
                          current: Employee = Depends(get_current_employee)):
    if entity_type not in ("product", "service", "bundle"):
        raise HTTPException(400, "entity_type must be product|service|bundle")
    return {"entity_type": entity_type,
            "tree": await _category_tree(db, current.org_id, entity_type, template_key)}


async def _sync_part_category_name(db, org_id, part):
    """Keep the denormalized Part.category name in step with category_id, so the POS
    and inventory (which group by the name string) reflect the managed tree."""
    if part.category_id:
        cat = (await db.execute(select(Category).where(
            Category.id == part.category_id, Category.org_id == org_id))).scalar_one_or_none()
        part.category = cat.name if cat else part.category
    else:
        part.category = None


class CategoryIn(BaseModel):
    entity_type: str = "product"
    name: str
    parent_id: Optional[str] = None
    auto_deduct: bool = False
    template_key: Optional[str] = None


@router.post("/categories", status_code=201)
async def create_category(data: CategoryIn, db: AsyncSession = Depends(get_db),
                          current: Employee = Depends(require_perm("can_manage_inventory"))):
    if data.entity_type not in ("product", "service", "bundle"):
        raise HTTPException(400, "entity_type must be product|service|bundle")
    name = (data.name or "").strip()
    if not name:
        raise HTTPException(400, "Category name is required")
    c = Category(org_id=current.org_id, entity_type=data.entity_type, name=name,
                 parent_id=data.parent_id or None, auto_deduct=data.auto_deduct,
                 template_key=(data.template_key or None))
    db.add(c); await db.commit(); await db.refresh(c)
    return {"id": c.id, "name": c.name, "entity_type": c.entity_type,
            "parent_id": c.parent_id, "template_key": c.template_key, "count": 0}


class CategoryPatch(BaseModel):
    name: Optional[str] = None
    parent_id: Optional[str] = None
    template_key: Optional[str] = None


@router.patch("/categories/{cid}")
async def update_category(cid: str, data: CategoryPatch, db: AsyncSession = Depends(get_db),
                          current: Employee = Depends(require_perm("can_manage_inventory"))):
    c = (await db.execute(select(Category).where(
        Category.id == cid, Category.org_id == current.org_id))).scalar_one_or_none()
    if not c:
        raise HTTPException(404, "Category not found")
    sent = data.model_fields_set
    if data.name is not None and data.name.strip():
        old = c.name
        c.name = data.name.strip()
        # keep product name-mirrors in sync after a rename
        parts = (await db.execute(select(Part).where(
            Part.org_id == current.org_id, Part.category_id == cid, Part.category == old))).scalars().all()
        for p in parts:
            p.category = c.name
    if "parent_id" in sent:
        if data.parent_id == cid:
            raise HTTPException(400, "A category can't be its own parent")
        # 144bc security: block moving a node under one of its own descendants (would orphan a subtree)
        if data.parent_id:
            all_cats = (await db.execute(select(Category).where(Category.org_id == current.org_id))).scalars().all()
            kids: dict = {}
            for x in all_cats:
                kids.setdefault(x.parent_id, []).append(x.id)
            stack, descendants = list(kids.get(cid, [])), set()
            while stack:
                nid = stack.pop()
                if nid in descendants:
                    continue
                descendants.add(nid)
                stack.extend(kids.get(nid, []))
            if data.parent_id in descendants:
                raise HTTPException(400, "Can't move a category under one of its own sub-categories")
        c.parent_id = data.parent_id or None
    if "template_key" in sent:
        c.template_key = data.template_key or None
    await db.commit()
    return {"id": c.id, "name": c.name, "parent_id": c.parent_id, "template_key": c.template_key}


@router.delete("/categories/{cid}")
async def delete_category(cid: str, db: AsyncSession = Depends(get_db),
                          current: Employee = Depends(require_perm("can_manage_inventory"))):
    c = (await db.execute(select(Category).where(Category.id == cid, Category.org_id == current.org_id))).scalar_one_or_none()
    if not c:
        raise HTTPException(404, "Category not found")
    # block delete if in use — the tree is shared across products, services AND bundles,
    # so all three must be checked or a service/bundle category delete would orphan rows.
    from backend.models import Service, Bundle
    kids = (await db.execute(select(func.count()).select_from(Category).where(Category.parent_id == cid))).scalar() or 0
    used = (await db.execute(select(func.count()).select_from(Part).where(Part.category_id == cid))).scalar() or 0
    used += (await db.execute(select(func.count()).select_from(Service).where(Service.category_id == cid))).scalar() or 0
    used += (await db.execute(select(func.count()).select_from(Bundle).where(Bundle.category_id == cid))).scalar() or 0
    if kids or used:
        raise HTTPException(409, "Category is in use (has sub-categories, products, services or bundles) — move them first")
    await db.delete(c); await db.commit()
    return {"ok": True}


async def _resolve_category_path(db, org_id, path: str) -> Optional[str]:
    """Resolve / create a 'A > B > C' product category path, return leaf id."""
    if not path:
        return None
    parent = None
    for name in [p.strip() for p in path.split(">") if p.strip()]:
        existing = (await db.execute(select(Category).where(
            Category.org_id == org_id, Category.entity_type == "product",
            Category.name == name, Category.parent_id == parent))).scalar_one_or_none()
        if not existing:
            existing = Category(org_id=org_id, entity_type="product", name=name, parent_id=parent)
            db.add(existing); await db.flush()
        parent = existing.id
    return parent


# ── Product master ───────────────────────────────────────────────────────────
def _ser_product(p: Part, qty=0.0, reserved=0.0, variant_count=0) -> dict:
    margin = round((p.sell_price - p.avg_cost) / p.avg_cost * 100, 1) if p.avg_cost else None
    return {"id": p.id, "name": p.name, "code": p.code, "sku": p.sku, "mpn": p.mpn, "barcode": p.barcode,
            "barcode_type": p.barcode_type, "brand": p.brand, "category_id": p.category_id,
            "description": p.description, "default_supplier_id": p.default_supplier_id,
            "uom": p.uom, "cost_price": p.cost_price, "avg_cost": p.avg_cost, "sell_price": p.sell_price,
            "min_stock": p.min_stock, "serialized": p.serialized, "expiry_tracking": p.expiry_tracking,
            "weight": p.weight, "dimensions": p.dimensions, "image_data": p.image_data, "warranty_days": p.warranty_days,
            "supplier_warranty_days": p.supplier_warranty_days, "parent_id": p.parent_id,
            "variant_attrs": p.variant_attrs, "variant_count": variant_count,
            "custom_data": p.custom_data or {},
            "compatibility_json": p.compatibility_json or {},
            "margin_pct": margin, "qty_available": qty, "qty_reserved": reserved,
            "low_stock": qty <= (p.min_stock or 0), "active": p.active}


async def _stock_for(db, part_id, warehouse_id=None, wh_location_id=None):
    q = select(StockLevel).where(StockLevel.part_id == part_id)
    if warehouse_id:
        # Match the warehouse's own rows, plus location-level rows (warehouse_id NULL) for the
        # same location — postings/sales write location-keyed stock, so a warehouse filter must
        # include those or it hides most on-hand (same split handled in the stocktake engine).
        cond = StockLevel.warehouse_id == warehouse_id
        if wh_location_id:
            cond = cond | ((StockLevel.warehouse_id.is_(None)) & (StockLevel.location_id == wh_location_id))
        q = q.where(cond)
    sls = (await db.execute(q)).scalars().all()
    return sum(s.qty_available for s in sls), sum(s.qty_reserved for s in sls)


@router.get("/products")
async def list_products(search: Optional[str] = None, category_id: Optional[str] = None,
                        warehouse_id: Optional[str] = None, include_variants: bool = False,
                        db: AsyncSession = Depends(get_db), current: Employee = Depends(get_current_employee)):
    q = select(Part).where(Part.org_id == current.org_id, Part.active == True)  # noqa: E712
    if not include_variants:
        q = q.where(Part.parent_id.is_(None))  # base products only
    if search:
        q = q.where(or_(*(_contains(c, search) for c in (Part.name, Part.sku, Part.barcode, Part.code, Part.mpn))))
    if category_id:
        q = q.where(Part.category_id == category_id)
    wh_location_id = None
    if warehouse_id:
        from backend.models import Warehouse
        wh_location_id = (await db.execute(select(Warehouse.location_id).where(
            Warehouse.id == warehouse_id, Warehouse.org_id == current.org_id))).scalar_one_or_none()
    rows = (await db.execute(q.order_by(Part.name))).scalars().all()
    # Stock and variant counts for the whole page in two grouped queries (this used to be two
    # queries per product — ~2,000 round trips for a 1,000-item catalogue).
    ids = [p.id for p in rows]
    stock: dict = {}
    variants: dict = {}
    if ids:
        sq = select(StockLevel.part_id, func.coalesce(func.sum(StockLevel.qty_available), 0),
                    func.coalesce(func.sum(StockLevel.qty_reserved), 0)).where(StockLevel.part_id.in_(ids))
        if warehouse_id:
            cond = StockLevel.warehouse_id == warehouse_id
            if wh_location_id:
                cond = cond | ((StockLevel.warehouse_id.is_(None)) & (StockLevel.location_id == wh_location_id))
            sq = sq.where(cond)
        stock = {pid: (qty, res) for pid, qty, res in (await db.execute(sq.group_by(StockLevel.part_id))).all()}
        variants = dict((await db.execute(select(Part.parent_id, func.count()).where(
            Part.parent_id.in_(ids)).group_by(Part.parent_id))).all())
    out = []
    for p in rows:
        # warehouse_id scopes the stock figure to one warehouse; parts with no stock there drop out.
        qty, res = stock.get(p.id, (0, 0))
        if warehouse_id and qty == 0 and res == 0:
            continue
        out.append(_ser_product(p, qty, res, variants.get(p.id, 0)))
    return {"items": out}


async def _checked_mpn(db, org_id, mpn, exclude_id=None) -> Optional[str]:
    """Normalise an MPN (trimmed, blank = none) and refuse one another active product of this
    org already uses — two cards with one MPN make scans and supplier imports ambiguous."""
    v = (mpn or "").strip()[:64] or None
    if v:
        q = select(Part).where(Part.org_id == org_id, Part.active == True,  # noqa: E712
                               func.lower(Part.mpn) == v.lower())
        if exclude_id:
            q = q.where(Part.id != exclude_id)
        clash = (await db.execute(q)).scalars().first()
        if clash:
            raise HTTPException(409, f"MPN {v} is already used by {clash.name}")
    return v


async def commit_mpn_safe(db) -> None:
    """Commit; a concurrent save that slipped past _checked_mpn hits uq_parts_org_mpn_active → 409."""
    from sqlalchemy.exc import IntegrityError
    try:
        await db.commit()
    except IntegrityError as e:
        await db.rollback()
        if "uq_parts_org_mpn_active" in str(e.orig):
            raise HTTPException(409, "That MPN is already used by another product")
        raise


class ProductIn(BaseModel):
    name: str
    code: Optional[str] = None
    sku: Optional[str] = None
    mpn: Optional[str] = Field(None, max_length=64)   # opa: manufacturer part number
    barcode: Optional[str] = None
    barcode_type: str = "code128"
    brand: Optional[str] = None
    description: Optional[str] = None
    default_supplier_id: Optional[str] = None
    category_id: Optional[str] = None
    uom: str = "pcs"
    cost_price: float = 0.0
    sell_price: float = 0.0
    min_stock: int = 0
    serialized: bool = False
    expiry_tracking: bool = False
    weight: Optional[float] = None
    dimensions: Optional[str] = None
    image_data: Optional[str] = None
    warranty_days: int = 0
    supplier_warranty_days: int = 0
    custom_data: Optional[dict] = None
    compatibility_json: Optional[dict] = None   # 144jt: {"fits": [{brand, model?}]}


@router.post("/products", status_code=201)
async def create_product(data: ProductIn, db: AsyncSession = Depends(get_db),
                         current: Employee = Depends(require_perm("can_manage_inventory"))):
    dup = (await db.execute(select(Part).where(Part.org_id == current.org_id, Part.name == data.name))).scalar_one_or_none()
    if dup:
        raise HTTPException(409, "A product with this name already exists (name is the unique identifier)")
    from backend.routers.forms import check_entity_required
    await check_entity_required(db, current.org_id, "product", data)   # built-ins the product form requires
    payload = data.model_dump()
    payload["mpn"] = await _checked_mpn(db, current.org_id, data.mpn)
    if not await can_see_costs(db, current):
        drop_cost_fields(payload)          # can't see costs, can't set them
    # enforce required custom fields (Settings → Forms → Products) on create
    payload["custom_data"] = await resolve_custom_data(
        db, current.org_id, "product", payload.get("custom_data"), creating=True)
    p = Part(org_id=current.org_id, **payload)
    await _sync_part_category_name(db, current.org_id, p)
    db.add(p); await commit_mpn_safe(db); await db.refresh(p)
    return _ser_product(p)


class BulkProductsIn(BaseModel):
    ids: list[str]
    action: str                       # category | price | delete
    category_id: Optional[str] = None
    price_pct: Optional[float] = None  # adjust by +/- percent
    price_set: Optional[float] = None  # set absolute price


@router.post("/products/bulk")
async def bulk_products(data: BulkProductsIn, db: AsyncSession = Depends(get_db),
                        current: Employee = Depends(require_perm("can_manage_inventory"))):
    """Build 144ap — bulk actions from the Products multi-select: move to a category,
    adjust price, or soft-delete (write-off keeps history)."""
    if not data.ids:
        raise HTTPException(422, "No products selected")
    if data.action == "delete":
        await assert_perm(db, current, "can_delete_products")
    parts = (await db.execute(select(Part).where(
        Part.id.in_(data.ids), Part.org_id == current.org_id))).scalars().all()
    n = 0
    blocked: list[str] = []
    for p in parts:
        if data.action == "category":
            p.category_id = data.category_id or None
            await _sync_part_category_name(db, current.org_id, p)
        elif data.action == "price":
            if data.price_set is not None:
                p.sell_price = q2(float(data.price_set))
            elif data.price_pct is not None:
                p.sell_price = q2((p.sell_price or 0) * (1 + float(data.price_pct) / 100))
        elif data.action == "delete":
            # 144co: never silently orphan a job kit — skip parts still in use, report them.
            from backend.services.catalog_links import catalog_usage
            if await catalog_usage(db, current.org_id, "product", p.id):
                blocked.append(p.name)
                continue
            p.active = False
        else:
            raise HTTPException(400, "Unknown action")
        n += 1
    await db.commit()
    return {"ok": True, "count": n, "blocked": blocked}


# ── Excel export / import ────────────────────────────────────────────────────
# MPN (opa) goes last so existing sheets keep their column positions; import matches by header.
_COLS = ["Code", "SKU", "Barcode", "Name", "Category", "UoM", "Cost price",
         "Sell price", "Min stock", "Client warranty (days)", "Serialized", "MPN"]


@router.get("/products/export")
async def export_products(db: AsyncSession = Depends(get_db), current: Employee = Depends(get_current_employee)):
    import openpyxl
    rows = (await db.execute(select(Part).where(Part.org_id == current.org_id, Part.parent_id.is_(None)))).scalars().all()
    # category id -> path
    cats = {c.id: c for c in (await db.execute(select(Category).where(
        Category.org_id == current.org_id, Category.entity_type == "product"))).scalars().all()}
    def path(cid):
        names = []
        while cid and cid in cats:
            names.append(cats[cid].name); cid = cats[cid].parent_id
        return " > ".join(reversed(names))
    show_cost = await can_see_costs(db, current)     # the cost column stays empty otherwise
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Products"; ws.append(_COLS)
    for p in rows:
        ws.append(safe_row([p.code, p.sku, p.barcode, p.name, path(p.category_id), p.uom, p.cost_price if show_cost else None,
                            p.sell_price, p.min_stock, p.warranty_days, "yes" if p.serialized else "no", p.mpn]))
    buf = io.BytesIO(); await asyncio.to_thread(wb.save, buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="products.xlsx"'})


@router.post("/products/import")
async def import_products(file: UploadFile = File(...), db: AsyncSession = Depends(get_db),
                          current: Employee = Depends(require_perm("can_manage_inventory"))):
    wb = open_xlsx(await file.read(MAX_XLSX_BYTES + 1))
    ws = wb.active
    if (ws.max_row or 0) - 1 > MAX_IMPORT_ROWS:
        raise HTTPException(413, too_many_rows(ws.max_row - 1))
    header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: header.index(name) for name in _COLS if name in header}
    if "Name" not in idx:
        raise HTTPException(400, "Import file must have a 'Name' column")
    created = updated = 0
    mpn_rows = []                       # (part, mpn, sheet row) — applied after every row is read
    for rownum, row in enumerate(ws.iter_rows(min_row=2), start=2):
        cells = [unguard_cell(c.value) for c in row]
        def g(col, cells=cells):
            i = idx.get(col)
            return cells[i] if i is not None and i < len(cells) else None
        name = g("Name")
        if not name:
            continue
        name = str(name).strip()
        cat_id = await _resolve_category_path(db, current.org_id, str(g("Category") or "").strip())
        existing = (await db.execute(select(Part).where(Part.org_id == current.org_id, Part.name == name))).scalar_one_or_none()
        fields = dict(code=str(g("Code")) if g("Code") else None, sku=str(g("SKU")) if g("SKU") else None,
                      barcode=str(g("Barcode")) if g("Barcode") else None, category_id=cat_id,
                      uom=str(g("UoM") or "pcs"), cost_price=float(g("Cost price") or 0),
                      sell_price=float(g("Sell price") or 0), min_stock=int(g("Min stock") or 0),
                      warranty_days=int(g("Client warranty (days)") or 0),
                      serialized=str(g("Serialized") or "no").strip().lower() in ("yes", "true", "1"))
        mpn = None
        if "MPN" in idx:          # an older sheet without the column leaves MPNs alone
            raw = g("MPN")
            mpn = str(raw).strip() if raw is not None else ""
            if len(mpn) > 64:     # same limit as the drawer / API (422 there)
                raise HTTPException(409, f"Row {rownum}: MPN is longer than 64 characters")
        if existing:
            target = existing
            for k, v in fields.items():
                if k == "serialized" and existing.serialized:
                    continue  # never un-serialize
                setattr(existing, k, v)
            updated += 1
        else:
            target = Part(org_id=current.org_id, name=name, **fields)
            db.add(target); created += 1
        if mpn is not None:
            mpn_rows.append((target, mpn or None, rownum))
    if mpn_rows:
        # Checked against the sheet's end state, so swapping two products' MPNs in one import
        # works: clear the touched MPNs first, then assign + check each (one active product per MPN).
        for part, _v, _n in mpn_rows:
            part.mpn = None
        await db.flush()
        seen: dict = {}
        for part, v, n in mpn_rows:
            if v:
                if v.lower() in seen:
                    raise HTTPException(409, f"Row {n}: MPN {v} is also on row {seen[v.lower()]}")
                seen[v.lower()] = n
                try:
                    await _checked_mpn(db, current.org_id, v, exclude_id=part.id)
                except HTTPException as e:
                    raise HTTPException(409, f"Row {n}: {e.detail}")
            part.mpn = v
    await commit_mpn_safe(db)
    return {"ok": True, "created": created, "updated": updated}


@router.get("/products/{pid}")
async def get_product(pid: str, db: AsyncSession = Depends(get_db), current: Employee = Depends(get_current_employee)):
    p = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(404, "Product not found")
    qty, res = await _stock_for(db, pid)
    variants = (await db.execute(select(Part).where(Part.parent_id == pid))).scalars().all()
    var_out = []
    for v in variants:
        vqty, vres = await _stock_for(db, v.id)
        var_out.append(_ser_product(v, vqty, vres))
    serials = []
    if p.serialized:
        srows = (await db.execute(select(SerializedItem).where(
            SerializedItem.part_id == pid, SerializedItem.status == "in_stock"))).scalars().all()
        serials = [{"id": s.id, "serial_number": s.serial_number, "location_id": s.location_id,
                    "bin_id": s.bin_id, "status": s.status} for s in srows]
    return {"product": _ser_product(p, qty, res, len(variants)), "variants": var_out, "serials_in_stock": serials}


@router.put("/products/{pid}")
async def update_product(pid: str, data: ProductIn, db: AsyncSession = Depends(get_db),
                         current: Employee = Depends(require_perm("can_manage_inventory"))):
    p = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(404, "Product not found")
    # A1: only fields the caller sent are written. A drawer that doesn't show weight, image or
    # barcode type must not reset them to the defaults (a QR product silently becoming Code128,
    # a product photo vanishing). The same goes for cost when it isn't sent.
    payload = data.model_dump(exclude_unset=True)
    payload["name"] = data.name
    if "mpn" in payload:
        payload["mpn"] = await _checked_mpn(db, current.org_id, data.mpn, exclude_id=p.id)
    if not await can_see_costs(db, current):
        drop_cost_fields(payload)
    # serialization is irreversible once enabled — cannot be turned off here
    if p.serialized and "serialized" in payload and not payload["serialized"]:
        raise HTTPException(409, "Serialization cannot be disabled once enabled (Orderry rule)")
    if payload.get("serialized") and not p.serialized:
        await assert_perm(db, current, "can_track_serials")
    # only touch custom_data when the caller explicitly sent it (avoid wiping on partial saves)
    if "custom_data" in data.model_fields_set:
        # enforce required + lock_after_create against the stored values. (144aw: this resolve was
        # mistakenly nested under the compatibility_json branch below, so required/lock went
        # unenforced on edits unless compatibility_json also happened to be sent.)
        payload["custom_data"] = await resolve_custom_data(
            db, current.org_id, "product", payload.get("custom_data"), existing=p.custom_data or {})
    for k, v in payload.items():
        setattr(p, k, v)
    if "category_id" in payload:
        await _sync_part_category_name(db, current.org_id, p)
    await commit_mpn_safe(db); await db.refresh(p)
    return _ser_product(p)


# ── Variants ─────────────────────────────────────────────────────────────────
class VariantIn(BaseModel):
    name: Optional[str] = None        # defaults to "<base> (<attrs>)"
    sku: Optional[str] = None
    barcode: Optional[str] = None
    sell_price: Optional[float] = None
    cost_price: Optional[float] = None
    variant_attrs: dict               # {"Size":"L","Color":"Red"}


@router.post("/products/{pid}/variants", status_code=201)
async def create_variant(pid: str, data: VariantIn, db: AsyncSession = Depends(get_db),
                         current: Employee = Depends(require_perm("can_manage_inventory"))):
    base = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not base:
        raise HTTPException(404, "Base product not found")
    if base.parent_id:
        raise HTTPException(409, "Cannot create a variant of a variant")
    label = ", ".join(f"{k}: {v}" for k, v in data.variant_attrs.items())
    name = data.name or f"{base.name} ({label})"
    v = Part(org_id=current.org_id, parent_id=base.id, name=name,
             sku=data.sku, barcode=data.barcode, category_id=base.category_id,
             uom=base.uom, barcode_type=base.barcode_type, serialized=base.serialized,
             cost_price=(data.cost_price if data.cost_price is not None and await can_see_costs(db, current)
                         else base.cost_price),
             sell_price=data.sell_price if data.sell_price is not None else base.sell_price,
             variant_attrs=data.variant_attrs)
    db.add(v); await db.commit(); await db.refresh(v)
    return _ser_product(v)


# ── Serialize toggle (irreversible) ──────────────────────────────────────────
@router.post("/products/{pid}/serialize")
async def enable_serialization(pid: str, db: AsyncSession = Depends(get_db),
                               current: Employee = Depends(require_perm("can_manage_inventory"))):
    p = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(404, "Product not found")
    if p.serialized:
        raise HTTPException(409, "Product is already serialized")
    await assert_perm(db, current, "can_track_serials")
    p.serialized = True
    await db.commit()
    return {"ok": True, "serialized": True,
            "note": "Serialization is now enabled and cannot be reverted. Every unit needs a serial number."}


# ── Barcode generator ────────────────────────────────────────────────────────
def _gen_barcode(btype: str) -> str:
    import random
    if btype == "ean13":
        body = "".join(random.choices("0123456789", k=12)); return body + gtin_check_digit(body)
    if btype == "ean8":
        body = "".join(random.choices("0123456789", k=7)); return body + gtin_check_digit(body)
    if btype == "upca":
        body = "".join(random.choices("0123456789", k=11)); return body + gtin_check_digit(body)
    # code128 / qr: free-form unique token
    return "RO" + "".join(random.choices("0123456789ABCDEFGHJKMNPQRSTUVWXYZ", k=10))


@router.post("/products/{pid}/barcode/generate")
async def generate_barcode(pid: str, type: str = Query("code128"),
                           db: AsyncSession = Depends(get_db),
                           current: Employee = Depends(require_perm("can_manage_inventory"))):
    if type not in ("upca", "ean13", "ean8", "code128", "qr"):
        raise HTTPException(400, "type must be upca|ean13|ean8|code128|qr")
    p = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(404, "Product not found")
    code = _gen_barcode(type)
    p.barcode = code; p.barcode_type = type
    await db.commit()
    return {"ok": True, "barcode": code, "barcode_type": type}


# ── Write-off (Orderry rules: serial vs non-serial, FEFO) ─────────────────────
class WriteOffIn(BaseModel):
    location_id: str
    qty: Optional[float] = None           # non-serialized
    serials: Optional[list[str]] = None   # serialized: specific serial numbers
    comment: Optional[str] = None


@router.post("/products/{pid}/write-off")
async def write_off(pid: str, data: WriteOffIn, db: AsyncSession = Depends(get_db),
                    current: Employee = Depends(require_perm("can_manage_inventory"))):
    p = (await db.execute(select(Part).where(Part.id == pid, Part.org_id == current.org_id))).scalar_one_or_none()
    if not p:
        raise HTTPException(404, "Product not found")
    sl = (await db.execute(select(StockLevel).where(
        StockLevel.part_id == pid, StockLevel.location_id == data.location_id))).scalar_one_or_none()

    if p.serialized:
        if not data.serials:
            raise HTTPException(400, "Serialized product: provide the serial numbers to write off")
        done = []
        for sn in data.serials:
            si = (await db.execute(select(SerializedItem).where(
                SerializedItem.part_id == pid, SerializedItem.serial_number == sn,
                SerializedItem.status == "in_stock"))).scalar_one_or_none()
            if not si:
                raise HTTPException(404, f"Serial '{sn}' not in stock")
            si.status = "written_off"
            done.append(sn)
        qty = len(done)
        if sl:
            sl.qty_available = max(0, sl.qty_available - qty)
        db.add(StockMovement(part_id=pid, location_id=data.location_id, movement_type="writeoff",
                             qty=qty, ref_type="writeoff", created_by_id=current.id))
        await db.commit()
        return {"ok": True, "written_off": done, "qty": qty}

    # non-serialized
    qty = data.qty or 0
    if qty <= 0:
        raise HTTPException(400, "Provide a positive qty to write off")
    if not sl or sl.qty_available < qty:
        raise HTTPException(409, f"Only {sl.qty_available if sl else 0:g} in stock at this location")
    # consume cost layers: FEFO when expiry tracking is on, else FIFO
    batches = (await db.execute(select(StockItemBatch).where(
        StockItemBatch.part_id == pid, StockItemBatch.location_id == data.location_id,
        StockItemBatch.qty_remaining > 0).order_by(
        StockItemBatch.expiry_date.asc().nullslast() if p.expiry_tracking
        else StockItemBatch.received_at.asc()))).scalars().all()
    remaining = qty
    for b in batches:
        if remaining <= 0:
            break
        take = min(float(b.qty_remaining), remaining)
        b.qty_remaining = float(b.qty_remaining) - take
        remaining -= take
    await _atomic_add(db, sl, "qty_available", -(qty), round_to=None)   # atomic
    db.add(StockMovement(part_id=pid, location_id=data.location_id, movement_type="writeoff",
                         qty=qty, ref_type="writeoff", created_by_id=current.id))
    await db.commit()
    return {"ok": True, "qty": qty, "method": "FEFO" if p.expiry_tracking else "FIFO",
            "qty_available": sl.qty_available}


async def _atomic_add(db, obj, attr, delta, *, round_to=None):
    from backend.services.money import add_to_column
    return await add_to_column(db, obj, attr, delta, round_to=round_to)
